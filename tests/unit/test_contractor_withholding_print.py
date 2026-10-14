"""Tests for printing تقرير ضرائب الخصم — نموذج 2 «الشريط الأخضر الداكن» (user, 2026-10-14).

* Each company on its own pages: نموذج 4's header and the filter strip on every page, the
  dark green band (إجمالي المستخلص، صافي الأعمال، رصيد أول المدة + ضرائب الخصم = الصافي)
  on its first page, then the table with a dark totals row.
* A company's رصيد أول المدة is the whole card of every contractor with a contract there,
  even when the card shows in two companies (user's choice).
* The text always wraps: the columns share the width (``table-layout:fixed``).
* «الإجمالي العام» after the last company: a line per company and the grand total.
* Excel: one right-to-left sheet per company with the band's figures over the table.
"""

from __future__ import annotations

import datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.reports.contracting_print import DATE, MONEY, PCT, PrintColumn, document_excel, document_html
from app.reports.contractor_withholding_print import (
    WithholdingDocument,
    WithholdingSection,
    grand_figures,
)
from app.services.contractor_withholding_report_service import ContractorWithholdingReportService

D = Decimal
COLUMNS = [PrintColumn("extract_date", "التاريخ", DATE), PrintColumn("contractor_name", "اسم المقاول"),
           PrintColumn("works_value", "إجمالي المستخلص", MONEY),
           PrintColumn("withholding_tax_pct", "نسبة ضرائب الخصم", PCT),
           PrintColumn("withholding_tax", "قيمة ضرائب الخصم", MONEY)]


def _row(day, name, works, pct, withholding):
    return {"extract_date": datetime.date(2026, 10, day), "contractor_name": name, "works_value": D(works),
            "withholding_tax_pct": D(pct), "withholding_tax": D(withholding)}


def _section(name, rows, opening):
    totals = {"works_value": sum((r["works_value"] for r in rows), D(0)),
              "before_tax": sum((r["works_value"] for r in rows), D(0)) * D("0.877"),
              "withholding_tax": sum((r["withholding_tax"] for r in rows), D(0)), "count": len(rows)}
    figures = {"works_value": totals["works_value"], "before_tax": totals["before_tax"], "opening": D(opening),
               "withholding_tax": totals["withholding_tax"], "net": D(opening) + totals["withholding_tax"]}
    return WithholdingSection({"company_name": name, "address": "قويسنا", "tax_registration_no": "300012533253262"},
                              rows, totals, figures=figures,
                              hints={"opening": "من بطاقات مقاولي الشركة", "works_value": f"{len(rows)} مستخلصات"})


def _sections():
    return [_section("شركة صابر", [_row(1, "أحمد محمد حافظ", "50000", "3", "1315.79"),
                                   _row(2, "محمود حسن", "150000", "1", "1315.79")], "41000"),
            _section("شركة درة", [_row(12, "نبيل غانم", "100000", "1", "877.19")], "0")]


def _document(sections, grand=True):
    return WithholdingDocument(title="تقرير ضرائب الخصم", subtitle="ضرائب الخصم على المستخلصات",
                               count_noun="مستخلصات", filters=[("الفترة", "كل الفترات")], columns=COLUMNS,
                               sections=sections, grand_totals=grand_figures(sections) if grand else None,
                               printed_by="admin")


# -- the figures ------------------------------------------------------------------------------


def test_the_grand_total_adds_up_the_companies():
    grand = grand_figures(_sections())
    assert grand["count"] == 3
    assert grand["opening"] == D("41000")
    assert grand["withholding_tax"] == D("3508.77")
    assert grand["net"] == D("44508.77")


# -- the page -----------------------------------------------------------------------------------


def test_each_company_has_its_header_band_and_table():
    page = document_html(_document(_sections()))  # the document brings its own html()
    assert page.count('class="section"') == 2
    assert page.count('class="band"') == 2  # one per company: its first page only (it is not ``.fixed``)
    assert page.count('class="head fixed"') == 2 and page.count('class="filters fixed"') == 2
    for caption in ("إجمالي المستخلص", "صافي الأعمال", "رصيد أول المدة", "ضرائب الخصم على المستخلصات",
                    "صافي ضرائب الخصم"):
        assert caption in page
    assert "41,000.00" in page and "43,631.58" in page  # شركة صابر: 41,000 + 2,631.58
    assert "من بطاقات مقاولي الشركة" in page
    assert 'class="s">+</div>' in page and 'class="s">=</div>' in page


def test_the_table_wraps_and_follows_the_columns():
    page = document_html(_document(_sections()))
    assert "table.wt { table-layout:fixed; }" in page  # «التفاف النص» always on in print (user)
    assert "<th>نسبة ضرائب الخصم</th>" in page and "<th>اسم الشركة</th>" not in page
    assert '<td class="num w">1,315.79</td>' in page  # the withholding in red, as on the screen
    assert "الإجمالي — 2 مستخلصات" in page and "2,631.58" in page


def test_the_grand_total_follows_the_last_company_only():
    page = document_html(_document(_sections()))
    assert page.count("الإجمالي العام — كل الشركات") == 1
    assert page.index("الإجمالي العام — كل الشركات") > page.index("شركة درة")
    assert "<th>الشركة</th>" in page and "44,508.77" in page
    single = document_html(_document(_sections()[:1], grand=False))
    assert "الإجمالي العام" not in single


def test_an_empty_company_says_so():
    page = document_html(_document([_section("شركة صابر", [], "41000")], grand=False))
    assert "لا توجد مستخلصات في الفترة" in page and "41,000.00" in page


def test_the_page_splitter_repeats_the_header_and_numbers_the_pages():
    page = document_html(_document(_sections()))
    assert "classList.contains('fixed')" in page and "classList.contains('split')" in page
    assert "page-no" in page and "' من '" in page


# -- Excel ------------------------------------------------------------------------------------


def test_excel_one_rtl_sheet_per_company_with_the_band_over_the_table(tmp_path):
    book = load_workbook(document_excel(_document(_sections()), tmp_path / "w.xlsx"))
    assert book.sheetnames == ["شركة صابر", "شركة درة"]
    sheet = book["شركة صابر"]
    assert sheet.sheet_view.rightToLeft
    values = [[cell for cell in row if cell is not None] for row in sheet.iter_rows(values_only=True)]
    flat = [value for row in values for value in row]
    assert "شركة صابر" in flat and "السجل الضريبي: 300012533253262" in flat and "تقرير ضرائب الخصم" in flat
    band = values.index(["إجمالي المستخلص", "صافي الأعمال", "رصيد أول المدة", "+ ضرائب الخصم على المستخلصات",
                         "= صافي ضرائب الخصم"])
    assert values[band + 1] == [200000, 175400, 41000, 2631.58, 43631.58]
    head = values.index(["التاريخ", "اسم المقاول", "إجمالي المستخلص", "نسبة ضرائب الخصم", "قيمة ضرائب الخصم"])
    assert head > band + 1
    assert ["الإجمالي — 2 مستخلصات", 200000, 2631.58] in values
    assert "الإجمالي العام — كل الشركات" not in flat
    last = [[cell for cell in row if cell is not None] for row in book["شركة درة"].iter_rows(values_only=True)]
    assert ["الإجمالي العام — كل الشركات"] in last
    assert ["شركة صابر", 2, 200000, 175400, 41000, 2631.58, 43631.58] in last
    assert last[-1][:2] == ["الإجمالي العام", 3]


# -- the service: each company's extracts and cards ---------------------------------------------


class _Db:
    def __init__(self, rows, cards):
        self.rows, self.cards, self.queries = rows, cards, []

    def fetch_one(self, sql, params=None):
        return {"amount": D(0)}

    def fetch_all(self, sql, params=None):
        self.queries.append((sql, params))
        if "contractor_extract_deductions" in sql:
            return []
        if "FROM client_companies" in sql:
            names = {2: ("A-H/CO-1001", "شركة صابر"), 9: ("A-H/CO-1006", "شركة درة"), 5: ("A-H/CO-1004", "شركة 5")}
            return [{"company_id": cid, "company_code": names[cid][0], "company_name": names[cid][1]}
                    for cid in sorted(params[0], key=lambda cid: names[cid][0])]
        if "GROUP BY w.company_id" in sql:
            return [{"company_id": cid, "amount": amount} for cid, amount in self.cards.items()]
        return [dict(row) for row in self.rows]


def _extract(extract_id, company_id, date, works, pct="1"):
    from app.services.contractor_extracts_report_service import report_row

    return report_row({
        "extract_id": extract_id, "extract_no": f"EX-{extract_id}", "extract_date": datetime.date(*date),
        "works_value": D(works), "vat_pct": D("14"), "advance_payment_pct": D("10"), "withholding_tax_pct": D(pct),
        "works_insurance_pct": D("5"), "social_insurance_pct": D("0"), "other_deductions": D("0"),
        "contractor_id": 4, "contractor_name": "أحمد", "project_id": company_id * 10, "project_name": "م",
        "company_id": company_id, "company_name": "ش", "contract_id": 1, "contract_no": "A-H/CT-1001",
        "contract_name": None})


def test_service_splits_the_report_by_company_and_repeats_a_card():
    # 456,000 → 400,000 → 1% = 4,000 in each company; the same contractor's card (36,000) in both
    db = _Db([_extract(1, 9, (2026, 3, 1), "456000"), _extract(2, 2, (2026, 1, 1), "456000")],
             {2: D("36000"), 9: D("36000"), 5: D("0")})
    parts = ContractorWithholdingReportService(db).withholding_by_company()
    assert [part["company"]["company_name"] for part in parts] == ["شركة صابر", "شركة درة"]  # by code; 5 left out
    for part in parts:
        assert part["report"]["opening"] == D("36000.00") and part["report"]["net"] == D("40000.00")
    card_sql, card_params = next(q for q in db.queries if "GROUP BY w.company_id" in q[0])
    assert "SELECT DISTINCT p.company_id, k.contractor_id" in card_sql and card_params == []


def test_service_moves_a_companys_older_extracts_into_its_own_opening():
    db = _Db([_extract(1, 9, (2026, 3, 1), "456000"), _extract(2, 2, (2026, 1, 1), "456000")], {2: D("36000")})
    parts = ContractorWithholdingReportService(db).withholding_by_company(date_from=datetime.date(2026, 2, 1))
    saber, dorra = (part["report"] for part in parts)
    assert saber["rows"] == [] and saber["opening"] == D("40000.00")
    assert len(dorra["rows"]) == 1 and dorra["opening"] == D("0.00") and dorra["net"] == D("4000.00")


def test_service_card_query_follows_the_filters():
    db = _Db([], {})
    ContractorWithholdingReportService(db).withholding_by_company(company_id=2, contractor_id=4, contract_id=7)
    card_sql, card_params = next(q for q in db.queries if "GROUP BY w.company_id" in q[0])
    assert "p.company_id = %s" in card_sql and "k.contractor_id = %s" in card_sql and "k.contract_id = %s" in card_sql
    assert card_params == [2, 4, 7]
