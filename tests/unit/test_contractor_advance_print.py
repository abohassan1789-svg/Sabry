"""Tests for printing تقرير كشف حساب الدفعة المقدمة — نموذج 3 (user, 2026-10-14).

* A block per contract: a band (the contract, its names, «مستخدم ٪»), the table of its
  extracts and دفعات مقدمة, and beside it the figures and the bar of what was paid back.
* «تحديد أعمدة الجدول» decides the band's names, the figures and the summary's columns.
* Each company prints on its own pages; after the blocks a summary table of its contracts,
  the last one ending with «الإجمالي العام — كل الشركات».
* The bar is coloured as on the screen, with a red mark when a contract used more than its advance.
* Excel: one right-to-left sheet per company with the blocks and then the summary table.
"""

from __future__ import annotations

import datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.reports.contracting_print import MONEY, PCT, SERIAL, PrintColumn, document_excel, document_html
from app.reports.contractor_advance_print import (
    REST_COLOR,
    AdvanceDocument,
    AdvanceSection,
    ContractBlock,
    band_parts,
    bar_parts,
    build_html,
    figure_cells,
    lines_caption,
    lines_totals,
    used_pct,
)

D = Decimal
ALL = ("contract_label", "contract_date", "company_name", "project_name", "contractor_name", "contract_value",
       "advance_payment_pct", "advance_agreed", "advance_deducted", "advance_paid_out", "remaining_advance")
COLUMNS = [SERIAL, PrintColumn("contract_label", "العقد"), PrintColumn("contract_value", "قيمة العقد", MONEY),
           PrintColumn("advance_payment_pct", "نسبة الدفعة المقدمة", PCT),
           PrintColumn("remaining_advance", "المتبقي", MONEY)]


def _contract(label, agreed, deducted, paid_out, value="1000000", pct="10"):
    agreed, deducted, paid_out = D(agreed), D(deducted), D(paid_out)
    return {"contract_label": label, "contract_date": datetime.date(2026, 9, 25), "contractor_name": "محمود حسن",
            "project_name": "مشروع 1", "company_name": "شركة صابر", "contract_value": D(value),
            "advance_payment_pct": D(pct), "advance_agreed": agreed, "advance_deducted": deducted,
            "advance_paid_out": paid_out, "remaining_advance": agreed - deducted - paid_out}


def _extract(no, works, advance):
    return {"extract_date": datetime.date(2026, 9, 25), "extract_no": no, "works_value": D(works),
            "advance_payment_pct": D("10"), "advance_payment": D(advance)}


def _payment(amount):
    return {"kind": "payment", "extract_date": datetime.date(2026, 10, 9), "extract_no": "دفعة مقدمة",
            "works_value": None, "advance_payment_pct": None, "advance_payment": D(amount)}


def _block(contract, lines):
    colors = ["#1D4ED8" if line.get("kind") == "payment" else "#137A38" for line in lines]
    return ContractBlock(contract, lines, colors)


def _document(sections, visible=ALL, grand=None):
    return AdvanceDocument("تقرير كشف حساب الدفعة المقدمة", "الدفعة المقدمة من كل عقد", [("الفترة", "كل الفترات")],
                           visible, COLUMNS, sections, grand, printed_by="admin",
                           printed_at=datetime.datetime(2026, 10, 14, 12, 0))


def _sections():
    first = AdvanceSection(
        {"company_id": 2, "company_name": "شركة صابر", "address": "قويسنا", "tax_registration_no": "300012533253262"},
        [_block(_contract("A-H/CT-1001", "100000", "30000", "0"),
                [_extract("A-H/CT-1001/EX-01", "100000", "10000"), _extract("A-H/CT-1001/EX-02", "200000", "20000")]),
         _block(_contract("A-H/CT-1007", "50000", "0", "10000"), [_payment("5000"), _payment("5000")])],
        {"contract_value": D("2000000"), "remaining_advance": D("110000")})
    second = AdvanceSection(
        {"company_id": 9, "company_name": "شركة درة"},
        [_block(_contract("A-H/CT-1009", "100000", "80000", "45000"),
                [_extract("A-H/CT-1009/EX-01", "800000", "80000"), _payment("45000")]),
         _block(_contract("A-H/CT-1011", "50000", "0", "0"), [])],
        {"contract_value": D("2000000"), "remaining_advance": D("25000")})
    return [first, second]


# -- the figures ------------------------------------------------------------------------------


def test_used_share_counts_the_extracts_and_the_payments():
    assert used_pct(_contract("x", "50000", "0", "10000")) == D(20)
    assert used_pct(_contract("x", "100000", "80000", "45000")) == D(125)
    assert used_pct(_contract("x", "0", "0", "0")) == D(0)


def test_the_bar_has_a_part_per_line_and_the_grey_rest():
    parts, over = bar_parts(_sections()[0].blocks[0])
    assert [round(share, 2) for share, _color in parts] == [0.1, 0.2, 0.7]
    assert parts[-1][1] == REST_COLOR and over is None


def test_an_overdrawn_contract_marks_where_its_advance_ends():
    parts, over = bar_parts(_sections()[1].blocks[0])  # 125,000 taken from an advance of 100,000
    assert REST_COLOR not in [color for _share, color in parts]
    assert [color for _share, color in parts] == ["#137A38", "#1D4ED8"]
    assert round(over, 2) == 0.8


def test_the_lines_total_and_its_caption():
    lines = [_extract("EX-01", "100000", "10000"), _payment("5000")]
    assert lines_totals(lines) == {"works_value": D("100000"), "advance_payment": D("15000")}
    assert lines_caption(lines) == "إجمالي المسدد — 1 مستخلص + 1 دفعة مقدمة"
    payments = [_payment("5000"), _payment("5000")]
    assert lines_caption(payments) == "إجمالي المسدد — 2 دفعات مقدمة"
    assert lines_totals(payments)["works_value"] is None  # blank, not «0.00»


def test_the_column_chooser_decides_the_band_and_the_figures():
    contract = _contract("A-H/CT-1001", "100000", "30000", "0")
    document = _document([], visible=ALL)
    assert band_parts(document, contract) == ["محمود حسن", "مشروع 1", "تاريخ العقد 2026-09-25"]
    assert [caption for caption, _value in figure_cells(document, contract)] == [
        "قيمة العقد", "الدفعة المقدمة (10%)", "المسدد من المستخلصات", "المصروف من الدفعات"]
    few = _document([], visible=("contract_label", "project_name", "advance_agreed", "remaining_advance"))
    assert band_parts(few, contract) == ["مشروع 1"]
    assert figure_cells(few, contract) == [("الدفعة المقدمة", D("100000"))]


# -- the page ---------------------------------------------------------------------------------


def test_a_block_per_contract_under_each_companys_header():
    page = build_html(_document(_sections()))
    assert page.count('class="blk"') == 4
    assert page.count('class="head fixed"') == 2 and page.count('class="filters fixed"') == 2
    assert page.count('class="legend fixed"') == 2
    assert "شركة صابر" in page and "شركة درة" in page and "السجل الضريبي: 300012533253262" in page
    assert "مستخدم 30%" in page and "مستخدم 20%" in page
    assert "A-H/CT-1001/EX-02" in page and "إجمالي المسدد — 2 مستخلصات" in page
    assert page.count('<tr class="pay">') == 3


def test_the_figures_print_beside_the_table_and_a_negative_rest_in_red():
    page = build_html(_document(_sections()))
    assert "الدفعة المقدمة (10%)" in page and "المصروف من الدفعات" in page
    assert '<b class="neg">-25,000.00</b>' in page
    assert "تجاوز الدفعة المقدمة بـ 25,000.00" in page
    assert 'class="over"' in page


def test_a_contract_without_lines_says_so():
    page = build_html(_document(_sections()))
    assert "لا توجد مستخلصات ولا دفعات مقدمة على هذا العقد حتى الآن." in page


def test_hidden_columns_leave_the_block():
    visible = ("contract_label", "advance_agreed", "remaining_advance")
    page = build_html(_document(_sections(), visible=visible))
    assert "محمود حسن" not in page and "تاريخ العقد 2026" not in page
    assert "قيمة العقد</div>" not in page and "المسدد من المستخلصات" not in page
    assert "الدفعة المقدمة<b>" in page


def test_the_summary_follows_each_companys_blocks_and_the_grand_total_ends_the_last():
    page = build_html(_document(_sections(), grand={"contract_value": D("4000000"), "remaining_advance": D("135000")}))
    assert page.count('class="split"') == 2 and page.count("ملخص العقود") == 2
    assert page.index("A-H/CT-1007") < page.index("ملخص العقود") < page.index("شركة درة")
    assert page.count("الإجمالي العام — كل الشركات") == 1
    assert page.rindex("الإجمالي العام") > page.rindex("A-H/CT-1011")
    assert "الإجمالي — 2 عقود" in page and "4,000,000.00" in page and "135,000.00" in page


def test_the_page_splitter_moves_blocks_whole_and_numbers_the_pages():
    page = document_html(_document(_sections()))  # the document brings its own html()
    assert "page-no" in page and "' من '" in page
    assert "classList.contains('split')" in page and "(تابع)" in page


# -- Excel ------------------------------------------------------------------------------------


def test_excel_one_rtl_sheet_per_company_with_the_blocks_then_the_summary(tmp_path):
    path = document_excel(_document(_sections(), grand={"contract_value": D("4000000"),
                                                         "remaining_advance": D("135000")}), tmp_path / "a.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["شركة صابر", "شركة درة"]
    sheet = book["شركة صابر"]
    assert sheet.sheet_view.rightToLeft
    values = [[cell for cell in row if cell is not None] for row in sheet.iter_rows(values_only=True)]
    flat = [value for row in values for value in row]
    assert "شركة صابر" in flat and "تقرير كشف حساب الدفعة المقدمة" in flat
    bands = [row[0] for row in values if row and str(row[0]).startswith("A-H/CT-")]
    assert bands[0] == "A-H/CT-1001   ·   محمود حسن   ·   مشروع 1   ·   تاريخ العقد 2026-09-25   ·   مستخدم 30%"
    assert ["التاريخ", "رقم المستخلص / البيان", "إجمالي المستخلص", "نسبة الدفعة المقدمة",
            "قيمة الدفعة المقدمة"] in values
    assert ["إجمالي المسدد — 2 مستخلصات", 300000, 30000] in values
    assert ["إجمالي المسدد — 2 دفعات مقدمة", 10000] in values
    assert flat.index("ملخص العقود") > flat.index("إجمالي المسدد — 2 دفعات مقدمة")
    assert "الإجمالي العام — كل الشركات" not in flat
    last = [value for row in book["شركة درة"].iter_rows(values_only=True) for value in row if value is not None]
    assert "الإجمالي العام — كل الشركات" in last
    assert "لا توجد مستخلصات ولا دفعات مقدمة على هذا العقد حتى الآن." in last
