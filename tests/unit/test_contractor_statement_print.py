"""Tests for printing تقرير كشف حساب مقاول (user, 2026-10-14).

* نموذج 2 «الحسابات التلاتة ورا بعض» and نموذج 5 «مع رسم الأرصدة» are two tabs of the preview.
* Each company prints on its own pages (the contractor's card opens the first one only),
  with «الإجمالي العام — كل الشركات» at the end.
* The الخصومات الأخرى are listed per type at the end of each company.
* Excel: one right-to-left sheet per company with the three statements one after the other.
"""

from __future__ import annotations

import datetime
from decimal import Decimal

from PySide6.QtWebEngineWidgets import QWebEngineView  # noqa: F401 - before the QApplication
from PySide6.QtWidgets import QApplication
import pytest
from openpyxl import load_workbook

from app.reports.contracting_print import DATE, MONEY, TEXT, PrintColumn, preview_dialog
from app.reports.contractor_statement_print import (
    CHARTS,
    DEDUCTION,
    EXTRACT,
    LAYOUTS,
    OPENING,
    PAYMENT,
    STACKED,
    TOTAL,
    AccountTable,
    StatementDocument,
    StatementRow,
    StatementSection,
    band_text,
    build_html,
    chart_points,
    export_excel,
    grand_accounts,
)

D = Decimal
COLUMNS = [PrintColumn("date", "التاريخ", DATE, nowrap=True), PrintColumn("extract_no", "رقم المستخلص"),
           PrintColumn("other_deductions", "خصومات أخرى", MONEY), PrintColumn("net", "صافي المستخلص", MONEY),
           PrintColumn("paid", "التحصيلات / الدفعات", MONEY), PrintColumn("balance", "الرصيد التراكمي", MONEY)]
TITLES = (("رصيد جاري", "الرصيد الجاري", "#1A7A3C"), ("تأمين أعمال", "تأمين الأعمال", "#0369A1"),
          ("تأمينات اجتماعية", "التأمينات الاجتماعية", "#7C3AED"))


def _account(key, title, color, opening, columns=COLUMNS, negative=False):
    balance = D("-500") if negative else opening + D("3052.28") - D("5000")
    totals = {"opening": opening, "held": D("3052.28"), "paid": D("5000"), "balance": balance,
              "extracts_count": 1, "payments_count": 1, "net": D("3052.28"), "other_deductions": D("1000")}
    rows = [
        StatementRow(OPENING, {"balance": opening}, "رصيد أول المدة — بطاقة المقاول"),
        StatementRow(EXTRACT, {"date": datetime.date(2026, 9, 25), "extract_no": "A-H/CT-1001/EX-01",
                               "other_deductions": D("1000"), "net": D("3052.28"), "balance": opening + D("3052.28")}),
        StatementRow(DEDUCTION, label="• أخرى — رصيد سابق", amount=D("1000")),
        StatementRow(PAYMENT, {"date": datetime.date(2026, 9, 26), "extract_no": "دفعة نقدي — عامة",
                               "paid": D("5000"), "balance": balance}),
        StatementRow(TOTAL, totals, "الإجمالي"),
    ]
    return AccountTable(key, title, color, list(columns), rows, totals, [float(opening), float(opening) + 3052.28,
                                                                         float(balance)])


def _section(name, opening, logo=None, negative=False):
    company = {"company_id": 1, "company_name": name, "address": "قويسنا", "tax_registration_no": "300012533253262",
               "logo": logo, "logo_mime": "image/png"}
    accounts = [_account(key, title, color, opening if index == 0 else D("0"), negative=negative and index == 0)
                for index, (key, title, color) in enumerate(TITLES)]
    return StatementSection(company, accounts, {"remaining_advance": D("970000")}, {"total": D("37403.51")},
                            {"types": [("أخرى", D("1000"))], "total": D("1000")})


def _document(layout=STACKED, two=True):
    sections = [_section("شركة صابر", D("9000000"))]
    if two:
        sections.append(_section("شركة درة", D("0"), negative=True))
    return StatementDocument(
        title="تقرير كشف حساب مقاول", subtitle="أحمد محمد حافظ", period="كل الفترات",
        filters=[("المشروع", "الكل"), ("العقد", "الكل")], sections=sections,
        grand=grand_accounts(sections) if two else None, printed_by="admin",
        printed_at=datetime.datetime(2026, 10, 14, 12, 0), layout=layout)


def test_two_print_models_as_tabs():
    assert [key for key, _caption in LAYOUTS] == [STACKED, CHARTS]
    assert "نموذج 2" in LAYOUTS[0][1] and "نموذج 5" in LAYOUTS[1][1]


def test_stacked_prints_the_three_statements_of_each_company_under_their_bands():
    page = build_html(_document())
    assert page.count('class="block"') == 6
    assert page.count('class="charts"') == 0
    assert "الرصيد الجاري" in page and "تأمين الأعمال" in page and "التأمينات الاجتماعية" in page
    assert band_text({"opening": D("9000000"), "held": D("3052.28"), "paid": D("5000"), "balance": D("8998052.28")}) \
        in page
    assert page.count("شركة صابر") >= 1 and page.count("شركة درة") >= 1
    assert "السجل الضريبي: 300012533253262" in page


def test_the_header_and_strip_repeat_and_the_pages_are_numbered():
    page = build_html(_document())
    assert page.count('class="head fixed"') == 2 and page.count('class="filters fixed"') == 2
    assert "page-no" in page and "(تابع)" in page  # a statement split across pages says so
    assert "متبقي الدفعة المقدمة" in page and "970,000.00" in page and "37,403.51" in page


def test_charts_model_adds_one_chart_per_account():
    page = build_html(_document(CHARTS))
    assert page.count('class="charts"') == 2 and page.count("<polyline") == 6
    assert page.count('class="block"') == 6  # the three statements still print (user)


def test_each_chart_has_its_own_scale():
    points = chart_points([100.0, 200.0, 150.0])
    assert points[0] == (10.0, 60.0) and points[1] == (150.0, 10.0) and points[2][1] == 35.0
    assert chart_points([5.0]) == [(150.0, 35.0)]
    assert chart_points([]) == []


def test_grand_total_only_at_the_end_and_adds_the_companies_up():
    document = _document()
    page = build_html(document)
    assert page.count("الإجمالي العام — كل الشركات") == 1
    assert page.index("الإجمالي العام") > page.rindex("شركة درة")
    assert document.grand[0].totals["opening"] == D("9000000")
    assert document.grand[0].totals["held"] == D("6104.56")
    assert "الإجمالي العام" not in build_html(_document(two=False))


def test_deduction_lines_and_the_summary_per_type():
    page = build_html(_document())
    assert '<td class="l" colspan="2">• أخرى — رصيد سابق</td><td class="num">1,000.00</td>' in page
    assert page.count("الخصومات الأخرى حسب النوع") == 2 and "إجمالي الخصومات" in page


def test_a_deduction_goes_under_the_net_when_its_column_is_hidden():
    document = _document(two=False)
    for account in document.sections[0].accounts:
        account.columns = [c for c in account.columns if c.key != "other_deductions"]
    assert '<td class="num">(1,000.00)</td>' in build_html(document)
    for account in document.sections[0].accounts:
        account.columns = [c for c in account.columns if c.key != "net"]
    assert "• أخرى — رصيد سابق  (1,000.00)" in build_html(document)


def test_opening_and_total_rows_and_negative_balances():
    page = build_html(_document())
    assert '<tr class="op"><td colspan="2">رصيد أول المدة — بطاقة المقاول</td>' in page
    assert '<tr class="pay">' in page and '<tr class="total">' in page
    assert '<td class="num neg">-500.00</td>' in page


def test_text_columns_print_blank_for_a_missing_value():
    document = _document(two=False)
    document.sections[0].accounts[0].columns.insert(1, PrintColumn("contract_label", "العقد", TEXT))
    page = build_html(document)
    assert "None" not in page


def test_excel_one_rtl_sheet_per_company_with_the_three_statements(tmp_path):
    path = export_excel(_document(), tmp_path / "statement.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["شركة صابر", "شركة درة"]
    first, last = book["شركة صابر"], book["شركة درة"]
    assert first.sheet_view.rightToLeft and last.sheet_view.rightToLeft
    values = [cell.value for row in first.iter_rows() for cell in row if cell.value is not None]
    assert first["A1"].value == "شركة صابر" and "تقرير كشف حساب مقاول" in values
    assert sum(1 for v in values if isinstance(v, str) and v.startswith(("الرصيد الجاري", "تأمين الأعمال",
                                                                         "التأمينات الاجتماعية"))) == 3
    assert 9000000 in values and "(1,000.00)" not in values and 1000 in values
    assert "الخصومات الأخرى حسب النوع" in values
    assert "الإجمالي العام — كل الشركات" not in values
    assert "الإجمالي العام — كل الشركات" in [cell.value for row in last.iter_rows() for cell in row]


@pytest.fixture(scope="module")
def qt_app():
    return QApplication.instance() or QApplication([])


def test_the_preview_switches_the_model_and_remembers_it(qt_app, monkeypatch):
    document = _document()
    picked = []
    dialog = preview_dialog(None, document, True, True, LAYOUTS, picked.append)
    try:
        assert dialog.layout_tabs is not None and dialog.layout_tabs.count() == 2
        assert dialog.layout_tabs.currentIndex() == 0
        dialog.layout_tabs.setCurrentIndex(1)
        assert document.layout == CHARTS and picked == [CHARTS]
        assert "polyline" in open(dialog.view.url().toLocalFile(), encoding="utf-8").read()
    finally:
        dialog.reject()
        dialog.deleteLater()
    plain = preview_dialog(None, document, True, True)
    assert plain.layout_tabs is None
    plain.reject()
    plain.deleteLater()


# -- the service: one statement per company --------------------------------------------------


class _Db:
    def __init__(self):
        self.queries = []

    def fetch_one(self, sql, params=None):
        self.queries.append((sql, params))
        return {"current_balance": D("1000"), "works_insurance_amount": D("20"),
                "social_insurance_amount": D("3"), "withholding_tax_amount": D("7")}

    def fetch_all(self, sql, params=None):
        self.queries.append((sql, params))
        extract = {"extract_no": "EX-01", "works_value": D("10000"), "vat_pct": D("14"),
                   "advance_payment_pct": D("10"), "withholding_tax_pct": D("1"), "works_insurance_pct": D("5"),
                   "social_insurance_pct": D("0"), "other_deductions": D("0"), "contract_no": "CT",
                   "contract_name": None, "contractor_name": "أحمد"}
        if "FROM client_companies" in sql:
            return [{"company_id": 9}, {"company_id": 2}]  # company-code order: درة before صابر
        if "FROM contractor_contracts k" in sql and "x.extract_id" not in sql:
            return [{"contract_id": 1, "contract_value": D("100000"), "advance_payment_pct": D("10"), "company_id": 2},
                    {"contract_id": 2, "contract_value": D("50000"), "advance_payment_pct": D("0"), "company_id": 9}]
        if "FROM contractor_extracts x" in sql:
            return [dict(extract, extract_id=1, contract_id=1, extract_date=datetime.date(2026, 9, 1), company_id=2),
                    dict(extract, extract_id=2, contract_id=2, extract_date=datetime.date(2026, 9, 2), company_id=9)]
        if "FROM contractor_payments y" in sql:
            return [{"payment_id": 1, "payment_date": datetime.date(2026, 9, 3), "amount": D("500"),
                     "account_type": None, "extract_no": None, "company_id": 2},
                    {"payment_id": 2, "payment_date": datetime.date(2026, 9, 4), "amount": D("200"),
                     "account_type": None, "extract_no": None, "company_id": None}]  # no project
        return []


def test_service_one_statement_per_company_adding_up_to_the_screens():
    from app.services.contractor_statement_report_service import ContractorStatementReportService

    service = ContractorStatementReportService(_Db())
    parts = service.statements_by_company(contractor_id=4)
    assert [company for company, _statement in parts] == [9, 2]
    first, second = parts[0][1], parts[1][1]
    # the card opens the first company only; the payment with no project goes there too
    assert first["accounts"]["رصيد جاري"]["opening"] == D("1000.00")
    assert second["accounts"]["رصيد جاري"]["opening"] == D("0.00")
    assert first["accounts"]["رصيد جاري"]["totals"]["payments_count"] == 1
    assert second["accounts"]["رصيد جاري"]["totals"]["payments_count"] == 1
    assert first["withholding"]["opening"] == D("7.00") and second["withholding"]["opening"] == D("0.00")
    assert first["contracts"]["contract_value"] == D("50000.00")
    whole = service.statement(contractor_id=4)
    for account in ("رصيد جاري", "تأمين أعمال", "تأمينات اجتماعية"):
        assert sum(statement["accounts"][account]["totals"]["balance"] for _c, statement in parts) == \
            whole["accounts"][account]["totals"]["balance"]


def test_service_queries_carry_the_company():
    from app.services.contractor_statement_report_service import ContractorStatementReportService

    db = _Db()
    ContractorStatementReportService(db).statement(contractor_id=4)
    for marker in ("FROM contractor_contracts k", "FROM contractor_extracts x", "FROM contractor_payments y"):
        sql = next(sql for sql, _params in db.queries if marker in sql)
        assert "p.company_id" in sql


def test_service_without_movements_has_no_company():
    from app.services.contractor_statement_report_service import ContractorStatementReportService

    class Empty(_Db):
        def fetch_all(self, sql, params=None):
            return []

    assert ContractorStatementReportService(Empty()).statements_by_company(contractor_id=4) == []
