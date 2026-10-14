"""Tests for the contracting reports' print / PDF / Excel (نموذج 4 «ثلاث خانات», user 2026-10-14).

* The header carries the company's name, address, السجل الضريبي and logo (الشركات والمشاريع).
* Each company prints on its own pages; the last table ends with «الإجمالي العام».
* The whole header repeats on every page and shows «صفحة س من ص».
* Excel: the same header above the same table, one right-to-left sheet per company.
"""

from __future__ import annotations

import datetime
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.reports.contracting_print import (
    DATE,
    MONEY,
    PCT,
    SERIAL,
    PrintColumn,
    PrintDocument,
    PrintSection,
    build_html,
    cell_text,
    export_excel,
    label_span,
    total_cells,
)

COLUMNS = [
    SERIAL,
    PrintColumn("contract_label", "العقد"),
    PrintColumn("contract_date", "تاريخ العقد", DATE),
    PrintColumn("contract_value", "قيمة العقد", MONEY),
    PrintColumn("advance_payment_pct", "نسبة المقدمة", PCT),
    PrintColumn("remaining_advance", "المتبقي من المقدمة", MONEY),
    PrintColumn("status", "الحالة"),
]


def _png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (40, 20), "#1A7A3C").save(buffer, format="PNG")
    return buffer.getvalue()


PNG = _png()


def _row(label, value, pct, remaining):
    return {"contract_label": label, "contract_date": datetime.date(2026, 10, 9), "contract_value": Decimal(value),
            "advance_payment_pct": Decimal(pct), "remaining_advance": Decimal(remaining), "status": "جاري"}


def _document(two_companies=True):
    sabry = {"company_id": 2, "company_name": "شركة صابر", "address": "الدقي، الجيزة",
             "tax_registration_no": "123-456", "logo": PNG, "logo_mime": "image/png"}
    dorra = {"company_id": 9, "company_name": "شركة درة", "address": None, "tax_registration_no": None,
             "logo": None, "logo_mime": None}
    sections = [PrintSection(sabry, [_row("A-H/CT-1001", "2000000", "50", "970000"),
                                     _row("A-H/CT-1002", "500000", "12.5", "-37500")],
                             {"contract_value": Decimal("2500000"), "remaining_advance": Decimal("932500")})]
    if two_companies:
        sections.append(PrintSection(dorra, [_row("A-H/CT-1009", "374319.92", "10", "1000")],
                                     {"contract_value": Decimal("374319.92"), "remaining_advance": Decimal("1000")}))
    return PrintDocument(
        title="تقرير عقود المقاولين", subtitle="قيمة العقود", count_noun="عقود",
        filters=[("الفترة", "كل الفترات"), ("المشروع", "الكل"), ("المقاول", "الكل")],
        columns=COLUMNS, sections=sections,
        grand_totals={"contract_value": Decimal("2874319.92"), "remaining_advance": Decimal("933500")}
        if two_companies else None,
        printed_by="admin", printed_at=datetime.datetime(2026, 10, 14, 10, 30),
    )


def test_cell_text():
    assert cell_text(COLUMNS[3], Decimal("1940000")) == "1,940,000.00"
    assert cell_text(COLUMNS[4], Decimal("12.50")) == "12.5%"
    assert cell_text(COLUMNS[2], datetime.date(2026, 9, 25)) == "2026-09-25"


def test_the_total_label_covers_the_leading_text_columns():
    assert label_span(COLUMNS) == 3  # م، العقد، التاريخ
    assert total_cells(COLUMNS, {"contract_value": Decimal("-5"), "remaining_advance": Decimal("1")}) == [
        ("-5.00", True), ("", False), ("1.00", False), ("", False)]


def test_the_header_carries_the_company_from_companies_and_projects():
    page = build_html(_document())
    assert "شركة صابر" in page and "العنوان: الدقي، الجيزة" in page and "السجل الضريبي: 123-456" in page
    assert "data:image/png;base64," in page
    assert "تقرير عقود المقاولين" in page and "admin" in page and "2026-10-14 10:30" in page
    assert page.count('class="page-no"') == 2  # one header per company; the script repeats it per page


def test_a_company_without_logo_or_address_prints_just_its_name():
    page = build_html(_document())
    dorra = page[page.index("شركة درة") - 300:]
    assert "<img" not in dorra[:400] and "العنوان" not in dorra[:600]


def test_each_company_has_its_totals_and_the_last_one_the_grand_total():
    page = build_html(_document())
    assert page.count('class="section"') == 2
    assert "الإجمالي — 2 عقود" in page and "الإجمالي — 1 عقود" in page
    assert page.count("الإجمالي العام — كل الشركات") == 1
    assert page.index("الإجمالي العام") > page.index("A-H/CT-1009")
    assert "الإجمالي العام" not in build_html(_document(two_companies=False))


def test_figures_never_break_and_negatives_are_red():
    page = build_html(_document())
    assert '<td class="num neg">-37,500.00</td>' in page
    assert '<td class="num">2,000,000.00</td>' in page
    assert "white-space:nowrap; direction:ltr" in page


def test_the_pages_are_split_by_the_page_script():
    page = build_html(_document())
    assert "querySelectorAll('.section')" in page and "' من ' + pages.length" in page


def test_the_filters_strip():
    page = build_html(_document())
    assert "<span>الفترة: </span><b>كل الفترات</b>" in page


def test_excel_one_rtl_sheet_per_company_with_the_header(tmp_path):
    path = export_excel(_document(), tmp_path / "report.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["شركة صابر", "شركة درة"]
    first = book["شركة صابر"]
    assert first.sheet_view.rightToLeft
    values = [cell for row in first.iter_rows(values_only=True) for cell in row if cell is not None]
    assert "شركة صابر" in values and "العنوان: الدقي، الجيزة" in values and "تقرير عقود المقاولين" in values
    assert any("الفترة: كل الفترات" in str(v) for v in values)
    assert len(first._images) == 1  # the logo
    header_row = next(r for r in first.iter_rows() if r[0].value == "م")
    assert [c.value for c in header_row] == [c.label for c in COLUMNS]
    first_line = first[header_row[0].row + 1]
    assert first_line[3].value == 2000000 and first_line[3].number_format == "#,##0.00"
    assert first_line[4].value == 50
    assert "الإجمالي العام — كل الشركات" not in values
    last_values = [cell for row in book["شركة درة"].iter_rows(values_only=True) for cell in row if cell is not None]
    assert "الإجمالي العام — كل الشركات" in last_values and 2874319.92 in last_values


def test_a_saved_excel_opens_straight_away(tmp_path, monkeypatch):
    from PySide6.QtGui import QDesktopServices
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

    from app.reports import contracting_print

    QApplication.instance() or QApplication([])
    target = tmp_path / "report.xlsx"
    opened = []
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: opened.append(url.toLocalFile()) or True)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: pytest.fail("no message when it opens"))
    contracting_print.save_excel(None, _document())
    assert target.exists() and opened == [str(target).replace("\\", "/")]


def test_the_pdf_opens_when_printing_finishes(monkeypatch):
    from PySide6.QtGui import QDesktopServices
    from PySide6.QtWidgets import QApplication

    from app.reports import contracting_print

    QApplication.instance() or QApplication([])
    opened = []
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: opened.append(url.toLocalFile()) or True)
    contracting_print._pdf_done(None, "C:/x/report.pdf", True)
    assert opened == ["C:/x/report.pdf"]


# -- grouped lines and the totals cards (تقرير المستخلصات, user 2026-10-14) -------------------


def _grouped_document():
    from app.reports.contracting_print import DEDUCTION, GROUP, ROW, SUBTOTAL, PrintLine

    document = _document(two_companies=False)
    columns = [SERIAL, PrintColumn("extract_no", "رقم المستخلص", nowrap=True),
               PrintColumn("other_deductions", "خصومات أخرى", MONEY), PrintColumn("net", "صافي المستخلص", MONEY)]
    row = {"extract_no": "A-H/CT-1001/EX-01", "other_deductions": Decimal("1000"), "net": Decimal("3052.28")}
    section = document.sections[0]
    section.rows = [row]
    section.totals = {"other_deductions": Decimal("1000"), "net": Decimal("3052.28"), "works_value": Decimal("10000")}
    section.lines = [PrintLine(GROUP, label="أحمد — 1 مستخلص"), PrintLine(ROW, row),
                     PrintLine(DEDUCTION, label="• جزاءات — تأخير", amount=Decimal("1000")),
                     PrintLine(SUBTOTAL, label="إجمالي أحمد", totals=section.totals)]
    document.columns = columns
    document.cards = (("works_value", "إجمالي المستخلص"), ("net", "صافي المستخلص"))
    return document


def test_group_titles_subtotals_and_deduction_lines():
    page = build_html(_grouped_document())
    assert '<tr class="group"><td colspan="4">أحمد — 1 مستخلص</td></tr>' in page
    assert '<tr class="ded"><td class="l" colspan="2">• جزاءات — تأخير</td><td class="num">1,000.00</td>' in page
    assert '<tr class="sub"><td colspan="2">إجمالي أحمد</td>' in page
    assert '<td class="nw">A-H/CT-1001/EX-01</td>' in page


def test_a_deduction_goes_under_the_net_when_its_column_is_hidden():
    from app.reports.contracting_print import DEDUCTION, PrintLine, deduction_cells

    document = _grouped_document()
    document.columns = [c for c in document.columns if c.key != "other_deductions"]
    text, cells = deduction_cells(document, PrintLine(DEDUCTION, label="• جزاءات", amount=Decimal("1000")))
    assert text == "• جزاءات" and cells == ["(1,000.00)"]
    document.columns = document.columns[:2]
    text, cells = deduction_cells(document, PrintLine(DEDUCTION, label="• جزاءات", amount=Decimal("1000")))
    assert text == "• جزاءات  (1,000.00)"


def test_the_cards_are_boxed_and_only_on_the_first_page():
    page = build_html(_grouped_document())
    assert '<div class="cards"' in page and "<b>10,000.00</b>" in page and '<div class=dark>صافي المستخلص' in page
    assert "if (cards && !first)" in page


def test_a_group_title_never_ends_a_page():
    assert "last.className === 'group'" in build_html(_grouped_document())


def test_excel_has_the_cards_groups_and_deductions(tmp_path):
    book = load_workbook(export_excel(_grouped_document(), tmp_path / "x.xlsx"))
    values = [cell for row in book.active.iter_rows(values_only=True) for cell in row if cell is not None]
    assert "إجمالي المستخلص" in values and 10000 in values
    assert "أحمد — 1 مستخلص" in values and "إجمالي أحمد" in values and "• جزاءات — تأخير" in values


def test_type_shrinks_when_many_columns_print():
    from app.reports.contracting_print import font_px

    assert font_px([SERIAL] * 14) == 11.5 and font_px([SERIAL] * 17) == 10.5 and font_px([SERIAL] * 19) == 9.5
