"""طباعة تقرير كشف حساب مقاول — نموذجان كتبويبين في المعاينة (user, 2026-10-14).

The user picked two of ten mockups and asked for both as tabs of the preview:

* **نموذج 2 «الحسابات التلاتة ورا بعض»** (``STACKED``): the three statements — الرصيد
  الجاري، تأمين الأعمال، التأمينات الاجتماعية — one after the other, each under a coloured
  band carrying its رصيد أول المدة / مستخلصات / دفعات / المتبقي.
* **نموذج 5 «مع رسم الأرصدة»** (``CHARTS``): the same, after a row of three small charts,
  one per نوع الحساب, each on its own scale (a 9-million current balance never flattens the
  insurance lines).

The user's rules:

* Each company prints on its own pages (its own statement: the contractor's card opens
  the first company's only), with «الإجمالي العام — كل الشركات» at the end.
* The header is نموذج 4's three cells (logo, name, address, السجل الضريبي from الشركات
  والمشاريع); it repeats on every page with «صفحة س من ص».
* Each account prints the columns its tab shows («تحديد أعمدة الجدول», per tab).
* The الخصومات الأخرى are listed per type at the end; their lines print under their
  extract when «إظهار تفاصيل الخصومات» is on.
* Excel: one right-to-left sheet per company holding the three statements one after the
  other, as نموذج 2 (user), with the same header.

The page, the preview, the printer and the PDF are :mod:`app.reports.contracting_print`'s;
the document hands them its own ``html()`` and ``excel()``.
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.reports.contracting_print import (
    CELL_RULE,
    INK,
    MONEY,
    MUTED,
    PAGE_HEIGHT_PX,
    PCT,
    RED,
    SLATE,
    TOTAL_RULE,
    PrintColumn,
    _STYLE,
    _add_logo,
    _decimal,
    _esc,
    _excel_value,
    _sheet_title,
    cell_text,
    company_lines,
    font_px,
    label_span,
    logo_uri,
)

STACKED, CHARTS = "stacked", "charts"
LAYOUTS = ((STACKED, "نموذج 2 — الحسابات التلاتة"), (CHARTS, "نموذج 5 — مع رسم الأرصدة"))

# kinds of table row
OPENING, EXTRACT, PAYMENT, DEDUCTION, TOTAL = "opening", "extract", "payment", "deduction", "total"
OPENING_BG, PAYMENT_BG, DEDUCTION_BG = "#EEF4F4", "#F3FAF5", "#FEF2F2"
FIGURES = (("opening", "أول المدة"), ("held", "مستخلصات"), ("paid", "دفعات"), ("balance", "المتبقي"))


@dataclass(frozen=True)
class StatementRow:
    kind: str
    values: dict[str, Any] = field(default_factory=dict)  # by column key; a missing key prints blank
    label: str = ""                                       # OPENING / TOTAL / DEDUCTION text
    amount: Any = None                                    # DEDUCTION


@dataclass
class AccountTable:
    """One نوع الحساب's statement: its band, its columns and rows, its chart points."""

    key: str
    title: str
    color: str
    columns: list[PrintColumn]
    rows: list[StatementRow]
    totals: dict[str, Any]  # opening, held, paid, balance, extracts_count, payments_count
    series: list[float] = field(default_factory=list)


@dataclass
class StatementSection:
    """One company's pages."""

    company: dict[str, Any]
    accounts: list[AccountTable]
    contracts: dict[str, Any]
    withholding: dict[str, Any]
    deductions: dict[str, Any]
    general: dict[str, Any] | None = None  # «دفعات عامة», with a contract chosen


@dataclass
class StatementDocument:
    title: str
    subtitle: str
    period: str
    filters: list[tuple[str, str]]
    sections: list[StatementSection]
    grand: list[AccountTable] | None = None  # each account's totals over every company
    printed_by: str = ""
    printed_at: datetime.datetime = field(default_factory=datetime.datetime.now)
    layout: str = STACKED

    @property
    def file_stem(self) -> str:
        return f"{self.title} - {self.subtitle} - {self.printed_at:%Y-%m-%d}".replace("/", "-")

    def html(self) -> str:
        return build_html(self)

    def excel(self, path: str | Path) -> Path:
        return export_excel(self, path)


def grand_accounts(sections: list[StatementSection]) -> list[AccountTable]:
    """Each نوع الحساب added up over the companies (only its totals are printed)."""
    grand = []
    for index, first in enumerate(sections[0].accounts):
        totals: dict[str, Any] = {}
        for section in sections:
            for key, value in section.accounts[index].totals.items():
                totals[key] = totals.get(key, 0) + value
        grand.append(AccountTable(first.key, first.title, first.color, [], [], totals))
    return grand


def _money(value: Any) -> str:
    return cell_text(PrintColumn("", "", MONEY), value)


def band_text(totals: dict[str, Any]) -> str:
    return " · ".join(f"{caption} {_money(totals.get(key))}" for key, caption in FIGURES)


def deduction_target(columns: list[PrintColumn]) -> str | None:
    """Where a deduction's amount goes: under «خصومات أخرى», else in parentheses under the net."""
    keys = [column.key for column in columns]
    return next((key for key in ("other_deductions", "net") if key in keys), None)


# -- HTML -------------------------------------------------------------------------------------

_EXTRA_STYLE = f"""
.block {{ display:flex; flex-direction:column; }}
.band {{ display:flex; justify-content:space-between; align-items:center; gap:12px; color:#FFFFFF;
         padding:4px 12px; font-size:12.5px; font-weight:700; border-radius:4px 4px 0 0; }}
.band span {{ font-size:11px; font-weight:400; direction:rtl; }}
tr.op td {{ background:{OPENING_BG}; font-weight:700; }}
tr.pay td {{ background:{PAYMENT_BG}; }}
.charts {{ display:grid; grid-template-columns:repeat(3, minmax(0, 1fr)); gap:8px; }}
.chart {{ border:1px solid {CELL_RULE}; border-radius:6px; padding:5px 10px 4px; display:flex; flex-direction:column; gap:2px; }}
.chart .h {{ display:flex; justify-content:space-between; align-items:baseline; font-size:12.5px; font-weight:700; }}
.chart .h b {{ font-size:14px; direction:ltr; }}
.chart .f {{ display:flex; justify-content:space-between; gap:6px; font-size:10px; color:{MUTED}; }}
.chart .f b {{ color:{INK}; direction:ltr; }}
.chart .none {{ height:64px; display:flex; align-items:center; justify-content:center; color:{MUTED}; font-size:11px; }}
.small {{ width:auto; min-width:360px; }}
.cap {{ font-size:12.5px; font-weight:700; }}
.note {{ font-size:11px; color:{SLATE}; border:1px dashed {CELL_RULE}; border-radius:5px; padding:5px 12px; }}
.grand-t th {{ background:{TOTAL_RULE}; }}
.grand-t td.k {{ font-weight:700; text-align:right; padding-right:10px; }}
"""

# Pages: the header and the filter strip (``.fixed``) repeat on every page; a statement
# (``.block``) is split between its rows and carries its band and table head onto the next
# page; anything else moves whole. Then «س من ص» on every page.
_PAGINATE = """
(function () {
  var LIMIT = %(limit)d;
  var pages = [];
  Array.prototype.slice.call(document.querySelectorAll('.section')).forEach(function (section) {
    var parts = Array.prototype.slice.call(section.children);
    var fixed = parts.filter(function (el) { return el.classList.contains('fixed'); });
    var flow = parts.filter(function (el) { return !el.classList.contains('fixed'); });
    var page, count;
    function newPage() {
      page = document.createElement('div');
      page.className = 'page';
      fixed.forEach(function (el) { page.appendChild(el.cloneNode(true)); });
      section.parentNode.insertBefore(page, section);
      pages.push(page);
      count = 0;  // flowing parts on the page
    }
    newPage();
    flow.forEach(function (el) {
      if (!el.classList.contains('block')) {
        page.appendChild(el);
        if (page.offsetHeight > LIMIT && count > 0) { page.removeChild(el); newPage(); page.appendChild(el); }
        count += 1;
        return;
      }
      var rows = Array.prototype.slice.call(el.querySelectorAll('tbody tr'));
      var shell = el.cloneNode(true);
      shell.querySelector('tbody').innerHTML = '';
      var part = shell.cloneNode(true), body = part.querySelector('tbody'), before = count;
      page.appendChild(part);
      count += 1;
      rows.forEach(function (row) {
        body.appendChild(row);
        if (page.offsetHeight <= LIMIT || (body.children.length === 1 && before === 0)) { return; }
        body.removeChild(row);
        var started = body.children.length > 0;
        if (!started) { page.removeChild(part); }
        newPage();
        part = shell.cloneNode(true);
        if (started) { part.querySelector('.band b').textContent += ' (تابع)'; }
        body = part.querySelector('tbody');
        page.appendChild(part);
        before = 0;
        count = 1;
        body.appendChild(row);
      });
    });
    section.parentNode.removeChild(section);
  });
  pages.forEach(function (page, index) {
    var cell = page.querySelector('.page-no');
    if (cell) { cell.textContent = (index + 1) + ' من ' + pages.length; }
  });
})();
"""


def _header_html(document: StatementDocument, section: StatementSection) -> str:
    company = section.company
    uri = logo_uri(company)
    logo = f'<img src="{uri}" alt="">' if uri else ""
    lines = "".join(f'<div class="co-line">{_esc(line)}</div>' for line in company_lines(company))
    return (
        '<div class="head fixed">'
        f'<div class="co">{logo}<div><div class="co-name">{_esc(company.get("company_name"))}</div>{lines}</div></div>'
        f'<div class="title"><div class="t">{_esc(document.title)}</div>'
        f'<div class="st">{_esc(document.subtitle)}</div></div>'
        '<div class="meta">'
        f'<div class="r1 l"><span class="c">تاريخ الطباعة</span><b>{document.printed_at:%Y-%m-%d %H:%M}</b></div>'
        f'<div class="r1"><span class="c">المستخدم</span><b>{_esc(document.printed_by or "—")}</b></div>'
        f'<div class="l"><span class="c">الفترة</span><b>{_esc(document.period)}</b></div>'
        '<div><span class="c">الصفحة</span><b class="page-no">1 من 1</b></div>'
        "</div></div>"
    )


def section_filters(document: StatementDocument, section: StatementSection) -> list[tuple[str, str]]:
    """The strip: the filters, then the section's advance left and its withholding total."""
    return list(document.filters) + [
        ("متبقي الدفعة المقدمة", _money(section.contracts.get("remaining_advance"))),
        ("إجمالي ضرائب الخصم", _money(section.withholding.get("total"))),
    ]


def _filters_html(document: StatementDocument, section: StatementSection) -> str:
    return '<div class="filters fixed">' + "".join(
        f"<div><span>{_esc(caption)}: </span><b>{_esc(value)}</b></div>"
        for caption, value in section_filters(document, section)
    ) + "</div>"


def chart_points(series: list[float], width: float = 300, height: float = 70) -> list[tuple[float, float]]:
    """The series on its own scale: x across the width, the lowest value at the bottom."""
    if not series:
        return []
    low, high = min(series), max(series)
    span = high - low
    count = len(series)
    points = []
    for index, value in enumerate(series):
        x = width / 2 if count == 1 else 10 + index * (width - 20) / (count - 1)
        y = (height / 2 if span == 0 else height - 10 - (value - low) / span * (height - 20))
        points.append((round(x, 1), round(y, 1)))
    return points


def _chart_html(account: AccountTable) -> str:
    points = chart_points(account.series)
    balance = account.totals.get("balance") or 0
    if len(points) > 1:
        line = " ".join(f"{x},{y}" for x, y in points)
        dots = "".join(f'<circle cx="{x}" cy="{y}" r="3" fill="{account.color}"></circle>'
                       for x, y in points) if len(points) <= 40 else ""
        drawing = (f'<svg viewBox="0 0 300 70" width="100%" height="64" style="direction:ltr">'
                   '<line x1="10" y1="60" x2="290" y2="60" stroke="#CBD5E1" stroke-width="1"></line>'
                   f'<polyline points="{line}" fill="none" stroke="{account.color}" stroke-width="2.5"></polyline>'
                   f"{dots}</svg>")
    else:
        drawing = '<div class="none">مفيش حركة في الفترة دي</div>'
    figures = "".join(f"<span>{caption} <b>{_money(account.totals.get(key))}</b></span>"
                      for key, caption in FIGURES[:3])
    color = RED if _decimal(balance) < 0 else account.color
    return (f'<div class="chart"><div class="h" style="color:{account.color}">{_esc(account.title)}'
            f'<b style="color:{color}">{_money(balance)}</b></div>{drawing}<div class="f">{figures}</div></div>')


def _number_cell(column: PrintColumn, value: Any, text: str) -> str:
    if column.kind not in (MONEY, PCT) or text == "":
        css = ' class="nw"' if column.nowrap else ""
        return f"<td{css}>{_esc(text)}</td>"
    negative = column.kind == MONEY and _decimal(value) < 0
    return f'<td class="num{" neg" if negative else ""}">{_esc(text)}</td>'


def _row_html(account: AccountTable, row: StatementRow) -> str:
    columns = account.columns
    span = min(label_span(columns), len(columns))
    if row.kind in (OPENING, TOTAL):
        cells = [f'<td colspan="{span}">{_esc(row.label)}</td>']
        for column in columns[span:]:
            value = row.values.get(column.key)
            shown = column.kind == MONEY and value is not None
            cells.append(_number_cell(column, value, _money(value) if shown else ""))
        return f'<tr class="{"op" if row.kind == OPENING else "total"}">' + "".join(cells) + "</tr>"
    if row.kind == DEDUCTION:
        target = deduction_target(columns)
        amount = _money(row.amount)
        shown = amount if target == "other_deductions" else f"({amount})"
        text = row.label if target else f"{row.label}  ({amount})"
        cells = "".join(f'<td class="num">{_esc(shown if column.key == target else "")}</td>'
                        for column in columns[span:])
        return f'<tr class="ded"><td class="l" colspan="{span}">{_esc(text)}</td>{cells}</tr>'
    cells = []
    for column in columns:
        value = row.values.get(column.key)
        cells.append(_number_cell(column, value, "" if value is None else cell_text(column, value)))
    css = ' class="pay"' if row.kind == PAYMENT else ""
    return f"<tr{css}>" + "".join(cells) + "</tr>"


def _block_html(account: AccountTable) -> str:
    columns = account.columns
    colgroup = "<colgroup>" + "".join(
        f'<col style="width:{column.width}px">' if column.width else "<col>" for column in columns) + "</colgroup>"
    head = "".join(f"<th>{_esc(column.label)}</th>" for column in columns)
    body = "".join(_row_html(account, row) for row in account.rows)
    return (f'<div class="block"><div class="band" style="background:{account.color}"><b>{_esc(account.title)}</b>'
            f"<span>{_esc(band_text(account.totals))}</span></div>"
            f'<table style="font-size:{font_px(columns)}px">{colgroup}<thead><tr>{head}</tr></thead>'
            f"<tbody>{body}</tbody></table></div>")


def deduction_rows(section: StatementSection) -> list[tuple[str, Any]]:
    return list(section.deductions.get("types") or [])


def _deductions_html(section: StatementSection) -> str:
    rows = deduction_rows(section)
    if not rows:
        return ""
    body = "".join(f'<tr><td>{_esc(kind)}</td><td class="num">{_money(amount)}</td></tr>' for kind, amount in rows)
    return ('<div><table class="small" style="font-size:11px"><thead><tr><th colspan="2">الخصومات الأخرى حسب النوع'
            '</th></tr><tr><th>النوع</th><th>المبلغ</th></tr></thead>'
            f'<tbody>{body}<tr class="total"><td>إجمالي الخصومات</td>'
            f'<td class="num">{_money(section.deductions.get("total"))}</td></tr></tbody></table></div>')


def general_text(general: dict[str, Any]) -> str:
    parts = [f"{kind} {_money(amount)}" for kind, amount in general.get("types") or []]
    return (f"دفعات عامة على مشروع العقد (مش داخلة في الرصيد): {general.get('count') or 0} دفعات"
            + (" — " + " · ".join(parts) if parts else "") + f" — الإجمالي {_money(general.get('total'))}")


def _grand_html(grand: list[AccountTable]) -> str:
    head = "".join(f"<th>{caption}</th>" for _key, caption in (("", "الحساب"),) + FIGURES)
    body = "".join(
        f'<tr><td class="k" style="color:{account.color}">{_esc(account.title)}</td>'
        + "".join(_number_cell(PrintColumn(key, "", MONEY), account.totals.get(key), _money(account.totals.get(key)))
                  for key, _caption in FIGURES) + "</tr>"
        for account in grand)
    return ('<div><div class="cap">الإجمالي العام — كل الشركات</div>'
            f'<table class="grand-t" style="font-size:12px"><thead><tr>{head}</tr></thead><tbody>{body}</tbody>'
            "</table></div>")


def _section_html(document: StatementDocument, section: StatementSection, last: bool) -> str:
    parts = [_header_html(document, section), _filters_html(document, section)]
    if document.layout == CHARTS:
        parts.append('<div class="charts">' + "".join(_chart_html(a) for a in section.accounts) + "</div>")
    parts += [_block_html(account) for account in section.accounts]
    parts.append(_deductions_html(section))
    if section.general is not None:
        parts.append(f'<div class="note">{_esc(general_text(section.general))}</div>')
    if last and document.grand:
        parts.append(_grand_html(document.grand))
    return '<div class="section">' + "".join(part for part in parts if part) + "</div>"


def build_html(document: StatementDocument) -> str:
    sections = "".join(_section_html(document, section, index == len(document.sections) - 1)
                       for index, section in enumerate(document.sections))
    script = _PAGINATE % {"limit": PAGE_HEIGHT_PX}
    return (
        '<!DOCTYPE html><html dir="rtl" lang="ar"><head><meta charset="utf-8">'
        f"<title>{_esc(document.title)}</title><style>{_STYLE}{_EXTRA_STYLE}</style></head>"
        f"<body>{sections}<script>{script}</script></body></html>"
    )


# -- Excel ------------------------------------------------------------------------------------


def export_excel(document: StatementDocument, path: str | Path) -> Path:
    """One right-to-left sheet per company: the header lines, then the three statements one
    after the other (نموذج 2), the deductions per type, and on the last sheet the grand total."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    thin, dark = Side(style="thin", color=CELL_RULE[1:]), Side(style="thin", color=TOTAL_RULE[1:])
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    total_grid = Border(left=dark, right=dark, top=dark, bottom=dark)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    head_fill = PatternFill("solid", fgColor=SLATE[1:])
    total_fill = PatternFill("solid", fgColor="E2E8F0")
    fills = {OPENING: PatternFill("solid", fgColor=OPENING_BG[1:]), PAYMENT: PatternFill("solid", fgColor=PAYMENT_BG[1:]),
             DEDUCTION: PatternFill("solid", fgColor=DEDUCTION_BG[1:]), TOTAL: total_fill}
    used: set[str] = set()

    for index, section in enumerate(document.sections):
        company = section.company
        sheet = workbook.create_sheet(_sheet_title(company.get("company_name"), used))
        sheet.sheet_view.rightToLeft = True
        width = max([len(account.columns) for account in section.accounts] + [5])
        first_col = 1
        if company.get("logo"):
            first_col = 2
            _add_logo(sheet, company)
        lines = [(company.get("company_name") or "", Font(bold=True, size=16)),
                 *((line, Font(size=11, color=SLATE[1:])) for line in company_lines(company)),
                 (document.title, Font(bold=True, size=14)),
                 (document.subtitle, Font(size=11, color=SLATE[1:])),
                 ("   |   ".join(f"{caption}: {value}" for caption, value in
                                 [("الفترة", document.period)] + section_filters(document, section)), Font(size=11)),
                 (f"تاريخ الطباعة: {document.printed_at:%Y-%m-%d %H:%M}   |   المستخدم: {document.printed_by or '—'}",
                  Font(size=11, color=SLATE[1:]))]
        row_no = 1
        for text, font in lines:
            cell = sheet.cell(row=row_no, column=first_col, value=text)
            cell.font, cell.alignment = font, right
            if width > first_col:
                sheet.merge_cells(start_row=row_no, start_column=first_col, end_row=row_no, end_column=width)
            row_no += 1
        row_no += 1

        for account in section.accounts:
            columns = account.columns
            count = len(columns)
            band = sheet.cell(row=row_no, column=1, value=f"{account.title}   —   {band_text(account.totals)}")
            for col in range(1, count + 1):
                sheet.cell(row=row_no, column=col).fill = PatternFill("solid", fgColor=account.color[1:])
            band.font, band.alignment = Font(bold=True, color="FFFFFF"), right
            if count > 1:
                sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=count)
            row_no += 1
            for col, column in enumerate(columns, start=1):
                cell = sheet.cell(row=row_no, column=col, value=column.label)
                cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, color="FFFFFF"), head_fill, center, grid
            sheet.row_dimensions[row_no].height = 32
            row_no += 1
            span = min(label_span(columns), count)
            for line in account.rows:
                _excel_row(sheet, row_no, account, line, span, fills, grid, total_grid, center, right)
                row_no += 1
            row_no += 1

        deductions = deduction_rows(section)
        if deductions:
            title = sheet.cell(row=row_no, column=1, value="الخصومات الأخرى حسب النوع")
            title.font = Font(bold=True)
            row_no += 1
            for kind, amount in deductions + [("إجمالي الخصومات", section.deductions.get("total"))]:
                name, figure = sheet.cell(row=row_no, column=1, value=kind), sheet.cell(row=row_no, column=2)
                _excel_value(figure, PrintColumn("", "", MONEY), amount)
                for cell in (name, figure):
                    cell.border, cell.alignment = grid, center
                if kind == "إجمالي الخصومات":
                    name.font = figure.font = Font(bold=True)
                    name.fill = figure.fill = total_fill
                row_no += 1
            row_no += 1
        if section.general is not None:
            sheet.cell(row=row_no, column=1, value=general_text(section.general)).alignment = right
            row_no += 2
        if index == len(document.sections) - 1 and document.grand:
            sheet.cell(row=row_no, column=1, value="الإجمالي العام — كل الشركات").font = Font(bold=True, size=12)
            row_no += 1
            for col, caption in enumerate(["الحساب"] + [caption for _key, caption in FIGURES], start=1):
                cell = sheet.cell(row=row_no, column=col, value=caption)
                cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, color="FFFFFF"), head_fill, center, grid
            row_no += 1
            for account in document.grand:
                name = sheet.cell(row=row_no, column=1, value=account.title)
                name.font, name.alignment, name.border = Font(bold=True, color=account.color[1:]), center, total_grid
                for col, (key, _caption) in enumerate(FIGURES, start=2):
                    cell = sheet.cell(row=row_no, column=col)
                    _excel_value(cell, PrintColumn(key, "", MONEY), account.totals.get(key))
                    cell.font, cell.alignment, cell.border = Font(bold=True), center, total_grid
                row_no += 1

        for col in range(1, width + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _excel_row(sheet: Any, row_no: int, account: AccountTable, line: StatementRow, span: int, fills: dict[str, Any],
               grid: Any, total_grid: Any, center: Any, right: Any) -> None:
    from openpyxl.styles import Font

    columns = account.columns
    bold = line.kind in (OPENING, TOTAL)
    if line.kind in (OPENING, TOTAL, DEDUCTION):
        sheet.cell(row=row_no, column=1, value=line.label)
    target = deduction_target(columns) if line.kind == DEDUCTION else None
    for col, column in enumerate(columns, start=1):
        cell = sheet.cell(row=row_no, column=col)
        if line.kind in (OPENING, TOTAL):
            value = line.values.get(column.key)
            if col > span and column.kind == MONEY and value is not None:
                _excel_value(cell, column, value)
        elif line.kind == DEDUCTION:
            if col == 1 and target is None:
                cell.value = f"{line.label}  ({_money(line.amount)})"
            if column.key == target == "other_deductions":
                _excel_value(cell, column, line.amount)
            elif column.key == target:
                cell.value = f"({_money(line.amount)})"
        else:
            value = line.values.get(column.key)
            if value is not None:
                _excel_value(cell, column, value)
        cell.alignment = center
        cell.border = total_grid if line.kind == TOTAL else grid
        if line.kind in fills:
            cell.fill = fills[line.kind]
        color = "991B1B" if line.kind == DEDUCTION else (
            RED[1:] if column.kind == MONEY and isinstance(cell.value, (int, float, Decimal)) and cell.value < 0 else None)
        if bold or color:
            cell.font = Font(bold=bold, color=color)
    if line.kind in (OPENING, TOTAL, DEDUCTION):
        if line.kind == DEDUCTION:
            sheet.cell(row=row_no, column=1).alignment = right
        if span > 1:
            sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=span)


__all__ = [
    "CHARTS", "DEDUCTION", "EXTRACT", "LAYOUTS", "OPENING", "PAYMENT", "STACKED", "TOTAL",
    "AccountTable", "StatementDocument", "StatementRow", "StatementSection",
    "band_text", "build_html", "chart_points", "deduction_target", "export_excel", "general_text", "grand_accounts",
    "section_filters",
]
