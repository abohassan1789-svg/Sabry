"""طباعة / معاينة / PDF / Excel لتقارير المقاولات — المرحلة الثالثة (user, 2026-10-14).

The user picked **نموذج 4 «ثلاث خانات»** of ten mockups, with clearly visible
table borders:

* The header is one bordered box in three cells: the company (logo, name, address,
  السجل الضريبي — from الشركات والمشاريع, ``client_companies``), the report title, and a
  2×2 grid: تاريخ الطباعة، المستخدم، العدد، الصفحة.
* Under it a strip with the filters, then the table: a slate header row, full
  cell borders, a totals row.

The user's rules:

* Each company prints on its own pages, with its own header and totals; when
  there is more than one, the last table ends with «الإجمالي العام».
* The printed columns are the ones «تحديد أعمدة الجدول» shows on the screen.
* The whole header repeats on every page. The rows are split into A4-landscape
  pages by a small script in the page itself (it measures the real rows, so a
  wrapped name never pushes a row off the paper), which also writes
  «صفحة س من ص».
* Excel carries the same header above the same table: one sheet per company,
  right-to-left, real numbers.
* A saved PDF or Excel file opens straight away (user, 2026-10-14).

One HTML document is previewed, printed and saved as PDF through Chromium
(``QWebEngineView``), as in :mod:`app.ui.screens.tawrid_statement_print`.
Nothing here reads the database: the screen hands over a :class:`PrintDocument`.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import os
import re
import tempfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

# A4 landscape with 10 mm margins: the printable area in CSS px (96 per inch).
MARGIN_MM = 10.0
PAGE_WIDTH_PX = 1047   # (297 − 20) mm
PAGE_HEIGHT_PX = 700   # (210 − 20) mm = 718 px, less a little room for rounding

INK, SLATE, RULE, CELL_RULE, TOTAL_RULE = "#0F172A", "#334155", "#CBD5E1", "#475569", "#1E293B"
TINT, TOTAL_BG, MUTED, RED = "#F1F5F9", "#E2E8F0", "#475569", "#B91C1C"
FONTS = "Tahoma, 'Segoe UI', Arial, sans-serif"

# kinds of column: how a value is written, and whether the totals row sums it
TEXT, DATE, MONEY, PCT = "text", "date", "money", "pct"


@dataclass(frozen=True)
class PrintColumn:
    key: str
    label: str
    kind: str = TEXT
    width: int | None = None  # px on paper; ``None`` shares the rest of the width


@dataclass
class PrintSection:
    """One company's pages: its header data, its lines and their totals."""

    company: dict[str, Any]          # company_name, address, tax_registration_no, logo, logo_mime
    rows: list[dict[str, Any]]
    totals: dict[str, Any]           # by column key (money columns only)


@dataclass
class PrintDocument:
    title: str
    subtitle: str
    count_noun: str                              # «عقود» — the count cell and the totals label
    filters: list[tuple[str, str]]               # (caption, chosen value), shown in the strip
    columns: list[PrintColumn]
    sections: list[PrintSection]
    grand_totals: dict[str, Any] | None = None   # set when there is more than one company
    printed_by: str = ""
    printed_at: datetime.datetime = field(default_factory=datetime.datetime.now)

    @property
    def file_stem(self) -> str:
        return f"{self.title} - {self.printed_at:%Y-%m-%d}"


SERIAL = PrintColumn("_serial", "م", TEXT, 30)


# -- values ---------------------------------------------------------------------------


def _decimal(value: Any) -> Decimal:
    return Decimal(str(value if value not in (None, "") else 0))


def cell_text(column: PrintColumn, value: Any) -> str:
    if column.kind == MONEY:
        return f"{_decimal(value):,.2f}"
    if column.kind == PCT:
        return f"{_decimal(value).normalize():f}%"
    if column.kind == DATE:
        return f"{value:%Y-%m-%d}" if isinstance(value, (datetime.date, datetime.datetime)) else str(value or "")
    return str(value if value is not None else "")


def _negative(column: PrintColumn, value: Any) -> bool:
    return column.kind == MONEY and _decimal(value) < 0


def label_span(columns: list[PrintColumn]) -> int:
    """How many leading columns «الإجمالي» covers: up to the first figure column."""
    span = 0
    for column in columns:
        if column.kind in (MONEY, PCT):
            break
        span += 1
    return max(span, 1)


def total_cells(columns: list[PrintColumn], totals: dict[str, Any]) -> list[tuple[str, bool]]:
    """The cells after «الإجمالي»: each money total, blank elsewhere."""
    cells = []
    for column in columns[label_span(columns):]:
        if column.kind == MONEY and column.key in totals:
            cells.append((cell_text(column, totals[column.key]), _negative(column, totals[column.key])))
        else:
            cells.append(("", False))
    return cells


def logo_uri(company: dict[str, Any]) -> str:
    logo = company.get("logo")
    if not logo:
        return ""
    return f"data:{company.get('logo_mime') or 'image/png'};base64,{base64.b64encode(bytes(logo)).decode('ascii')}"


def company_lines(company: dict[str, Any]) -> list[str]:
    lines = []
    if company.get("address"):
        lines.append(f"العنوان: {company['address']}")
    if company.get("tax_registration_no"):
        lines.append(f"السجل الضريبي: {company['tax_registration_no']}")
    return lines


# -- HTML -------------------------------------------------------------------------------


def _esc(value: Any) -> str:
    return html.escape("" if value is None else str(value))


_STYLE = f"""
* {{ box-sizing:border-box; }}
html, body {{ margin:0; padding:0; }}
body {{ font-family:{FONTS}; color:{INK}; background:#E9E9EA;
       -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
.page, .section {{ width:{PAGE_WIDTH_PX}px; background:#FFFFFF; display:flex; flex-direction:column; gap:8px; }}
.page + .page {{ break-before:page; page-break-before:always; }}
@media screen {{ .page, .section {{ margin:26px auto; outline:16px solid #FFFFFF;
                  box-shadow:0 0 0 17px #D4D4D8; }} }}
@media print {{ body {{ background:#FFFFFF; }} .page, .section {{ margin:0; }} }}
.head {{ display:grid; grid-template-columns:repeat(3, minmax(0, 1fr)); border:1.5px solid {SLATE}; border-radius:6px; }}
.co {{ display:flex; align-items:center; gap:12px; padding:10px 14px; border-left:1.5px solid {SLATE}; min-width:0; }}
.co img {{ width:70px; height:70px; object-fit:contain; flex:none; }}
.co-name {{ font-size:19px; font-weight:700; overflow-wrap:anywhere; }}
.co-line {{ font-size:11.5px; color:{SLATE}; overflow-wrap:anywhere; }}
.title {{ display:flex; flex-direction:column; align-items:center; justify-content:center; gap:4px; padding:10px;
          background:{TINT}; text-align:center; }}
.title .t {{ font-size:22px; font-weight:700; }}
.title .st {{ font-size:12px; color:{SLATE}; }}
.meta {{ display:grid; grid-template-columns:repeat(2, minmax(0, 1fr)); border-right:1.5px solid {SLATE}; font-size:12px; }}
.meta div {{ padding:7px 12px; }}
.meta .c {{ color:{MUTED}; }}
.meta b {{ display:block; }}
.meta .r1 {{ border-bottom:1px solid {RULE}; }}
.meta .l {{ border-left:1px solid {RULE}; }}
.filters {{ display:flex; border:1px solid {RULE}; border-radius:6px; font-size:12px; }}
.filters div {{ flex:1; padding:6px 12px; border-left:1px solid {RULE}; overflow-wrap:anywhere; }}
.filters div:last-child {{ border-left:none; }}
.filters span {{ color:{MUTED}; }}
table {{ width:100%; border-collapse:collapse; font-size:11.5px; }}
th {{ background:{SLATE}; color:#FFFFFF; border:1px solid {SLATE}; padding:6px 3px; font-weight:700; }}
td {{ border:1px solid {CELL_RULE}; padding:6px 3px; text-align:center; overflow-wrap:break-word; }}
td.num {{ white-space:nowrap; direction:ltr; }}  /* a figure never breaks; its minus stays in front */
td.neg {{ color:{RED}; }}
tr.total td {{ background:{TOTAL_BG}; border-color:{TOTAL_RULE}; font-weight:700; }}
tr.grand td {{ background:{SLATE}; color:#FFFFFF; border-color:{TOTAL_RULE}; font-weight:700; }}
tr.grand td.neg {{ color:#FECACA; }}
td.empty {{ color:{MUTED}; padding:14px 3px; }}
"""

# Splits each section's rows into pages that fit the paper; every page repeats the
# section's header, filter strip and table head. Then writes «س من ص» on every page.
_PAGINATE = """
(function () {
  var LIMIT = %(limit)d;
  var pages = [];
  var sections = Array.prototype.slice.call(document.querySelectorAll('.section'));
  sections.forEach(function (section) {
    var rows = Array.prototype.slice.call(section.querySelectorAll('tbody tr'));
    var blank = section.cloneNode(true);
    blank.querySelector('tbody').innerHTML = '';
    var page, body;
    function newPage() {
      page = blank.cloneNode(true);
      page.className = 'page';
      body = page.querySelector('tbody');
      section.parentNode.insertBefore(page, section);
      pages.push(page);
    }
    newPage();
    rows.forEach(function (row) {
      body.appendChild(row);
      if (page.offsetHeight > LIMIT && body.children.length > 1) {
        body.removeChild(row);
        newPage();
        body.appendChild(row);
      }
    });
    section.parentNode.removeChild(section);
  });
  pages.forEach(function (page, index) {
    var cell = page.querySelector('.page-no');
    if (cell) { cell.textContent = (index + 1) + ' من ' + pages.length; }
  });
})();
"""


def _cell_class(column: PrintColumn, value: Any) -> str:
    if column.kind not in (MONEY, PCT):
        return ""
    return ' class="num neg"' if _negative(column, value) else ' class="num"'


def _colgroup(columns: list[PrintColumn]) -> str:
    return "<colgroup>" + "".join(
        f'<col style="width:{column.width}px">' if column.width else "<col>" for column in columns
    ) + "</colgroup>"


def _header_html(document: PrintDocument, section: PrintSection) -> str:
    company = section.company
    uri = logo_uri(company)
    logo = f'<img src="{uri}" alt="">' if uri else ""
    lines = "".join(f'<div class="co-line">{_esc(line)}</div>' for line in company_lines(company))
    return (
        '<div class="head">'
        f'<div class="co">{logo}<div><div class="co-name">{_esc(company.get("company_name"))}</div>{lines}</div></div>'
        f'<div class="title"><div class="t">{_esc(document.title)}</div>'
        f'<div class="st">{_esc(document.subtitle)}</div></div>'
        '<div class="meta">'
        f'<div class="r1 l"><span class="c">تاريخ الطباعة</span><b>{document.printed_at:%Y-%m-%d %H:%M}</b></div>'
        f'<div class="r1"><span class="c">المستخدم</span><b>{_esc(document.printed_by or "—")}</b></div>'
        f'<div class="l"><span class="c">عدد ال{_esc(document.count_noun)}</span><b>{len(section.rows)}</b></div>'
        '<div><span class="c">الصفحة</span><b class="page-no">1 من 1</b></div>'
        "</div></div>"
    )


def _filters_html(document: PrintDocument) -> str:
    return '<div class="filters">' + "".join(
        f"<div><span>{_esc(caption)}: </span><b>{_esc(value)}</b></div>" for caption, value in document.filters
    ) + "</div>"


def _total_row(document: PrintDocument, label: str, totals: dict[str, Any], css: str) -> str:
    span = label_span(document.columns)
    cells = "".join(f'<td class="num{" neg" if negative else ""}">{_esc(text)}</td>'
                    for text, negative in total_cells(document.columns, totals))
    return f'<tr class="{css}"><td colspan="{span}">{_esc(label)}</td>{cells}</tr>'


def _section_html(document: PrintDocument, section: PrintSection, last: bool) -> str:
    columns = document.columns
    head = "".join(f"<th>{_esc(column.label)}</th>" for column in columns)
    body = []
    for index, row in enumerate(section.rows, start=1):
        cells = []
        for column in columns:
            value = index if column is SERIAL else row.get(column.key)
            cells.append(f"<td{_cell_class(column, value)}>{_esc(cell_text(column, value))}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    if section.rows:
        body.append(_total_row(document, f"الإجمالي — {len(section.rows)} {document.count_noun}",
                               section.totals, "total"))
    else:
        body.append(f'<tr><td class="empty" colspan="{len(columns)}">لا توجد {_esc(document.count_noun)}</td></tr>')
    if last and document.grand_totals is not None:
        body.append(_total_row(document, "الإجمالي العام — كل الشركات", document.grand_totals, "grand"))
    return (
        '<div class="section">'
        f"{_header_html(document, section)}{_filters_html(document)}"
        f"<table>{_colgroup(columns)}<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"
        "</div>"
    )


def build_html(document: PrintDocument) -> str:
    sections = "".join(_section_html(document, section, index == len(document.sections) - 1)
                       for index, section in enumerate(document.sections))
    script = _PAGINATE % {"limit": PAGE_HEIGHT_PX}
    return (
        '<!DOCTYPE html><html dir="rtl" lang="ar"><head><meta charset="utf-8">'
        f"<title>{_esc(document.title)}</title><style>{_STYLE}</style></head>"
        f"<body>{sections}<script>{script}</script></body></html>"
    )


# -- Excel ------------------------------------------------------------------------------


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", " ", name or "تقرير").strip()[:28] or "تقرير"
    title, number = base, 2
    while title in used:
        title = f"{base[:25]} ({number})"
        number += 1
    used.add(title)
    return title


def export_excel(document: PrintDocument, path: str | Path) -> Path:
    """One right-to-left sheet per company: the header lines, then the table and its totals."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    workbook.remove(workbook.active)
    columns = document.columns
    width = len(columns)
    thin, thick = Side(style="thin", color=CELL_RULE[1:]), Side(style="thin", color=TOTAL_RULE[1:])
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    total_grid = Border(left=thick, right=thick, top=thick, bottom=thick)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head_fill = PatternFill("solid", fgColor=SLATE[1:])
    total_fill = PatternFill("solid", fgColor=TOTAL_BG[1:])
    grand_fill = PatternFill("solid", fgColor=SLATE[1:])
    used: set[str] = set()

    for index, section in enumerate(document.sections):
        sheet = workbook.create_sheet(_sheet_title(section.company.get("company_name"), used))
        sheet.sheet_view.rightToLeft = True
        company = section.company
        first_col = 1
        if company.get("logo"):
            first_col = 2  # the logo sits in column A, beside the company lines
            _add_logo(sheet, company)
        lines = [(company.get("company_name") or "", Font(bold=True, size=16)),
                 *((line, Font(size=11, color=SLATE[1:])) for line in company_lines(company)),
                 (document.title, Font(bold=True, size=14)),
                 (document.subtitle, Font(size=11, color=SLATE[1:])),
                 ("   |   ".join(f"{caption}: {value}" for caption, value in document.filters), Font(size=11)),
                 (f"تاريخ الطباعة: {document.printed_at:%Y-%m-%d %H:%M}   |   المستخدم: {document.printed_by or '—'}"
                  f"   |   عدد ال{document.count_noun}: {len(section.rows)}", Font(size=11, color=SLATE[1:]))]
        row_no = 1
        for text, font in lines:
            cell = sheet.cell(row=row_no, column=first_col, value=text)
            cell.font = font
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if width > first_col:
                sheet.merge_cells(start_row=row_no, start_column=first_col, end_row=row_no, end_column=width)
            row_no += 1
        row_no += 1

        for col, column in enumerate(columns, start=1):
            cell = sheet.cell(row=row_no, column=col, value=column.label)
            cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, color="FFFFFF"), head_fill, center, grid
        sheet.row_dimensions[row_no].height = 32
        row_no += 1
        for serial, row in enumerate(section.rows, start=1):
            for col, column in enumerate(columns, start=1):
                value = serial if column is SERIAL else row.get(column.key)
                cell = sheet.cell(row=row_no, column=col)
                _excel_value(cell, column, value)
                cell.alignment, cell.border = center, grid
                if _negative(column, value):
                    cell.font = Font(color=RED[1:])
            row_no += 1
        if section.rows:
            row_no = _excel_total(sheet, row_no, document, f"الإجمالي — {len(section.rows)} {document.count_noun}",
                                  section.totals, total_fill, Font(bold=True), total_grid, center)
        else:
            sheet.cell(row=row_no, column=1, value=f"لا توجد {document.count_noun}").alignment = center
            sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=width)
            row_no += 1
        if index == len(document.sections) - 1 and document.grand_totals is not None:
            _excel_total(sheet, row_no, document, "الإجمالي العام — كل الشركات", document.grand_totals,
                         grand_fill, Font(bold=True, color="FFFFFF"), total_grid, center)

        from openpyxl.utils import get_column_letter
        for col, column in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = max(6, (column.width or 110) / 7)
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _excel_value(cell: Any, column: PrintColumn, value: Any) -> None:
    if column.kind == MONEY:
        cell.value = float(_decimal(value))
        cell.number_format = "#,##0.00"
    elif column.kind == PCT:
        cell.value = float(_decimal(value))
        cell.number_format = '0.##"%"'
    else:
        cell.value = cell_text(column, value) if column.kind == DATE else value


def _excel_total(sheet: Any, row_no: int, document: PrintDocument, label: str, totals: dict[str, Any],
                 fill: Any, font: Any, border: Any, alignment: Any) -> int:
    span = label_span(document.columns)
    sheet.cell(row=row_no, column=1, value=label)
    for col, column in enumerate(document.columns, start=1):
        cell = sheet.cell(row=row_no, column=col)
        if col > span and column.kind == MONEY and column.key in totals:
            _excel_value(cell, column, totals[column.key])
        cell.fill, cell.font, cell.border, cell.alignment = fill, font, border, alignment
    if span > 1:
        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=span)
    return row_no + 1


def _add_logo(sheet: Any, company: dict[str, Any]) -> None:
    try:
        from openpyxl.drawing.image import Image
        from PIL import Image as PILImage

        picture = PILImage.open(io.BytesIO(bytes(company["logo"])))
        buffer = io.BytesIO()
        picture.convert("RGBA").save(buffer, format="PNG")
        buffer.seek(0)
        image = Image(buffer)
        scale = 70 / max(image.height, 1)
        image.width, image.height = image.width * scale, image.height * scale
        sheet.add_image(image, "A1")
        sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 12)
    except Exception:  # noqa: BLE001 - a bad logo never stops the export
        pass


# -- Qt: preview, print, PDF ------------------------------------------------------------


def page_layout() -> Any:
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize

    return QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Landscape,
                       QMarginsF(MARGIN_MM, MARGIN_MM, MARGIN_MM, MARGIN_MM), QPageLayout.Millimeter)


def _write_temp_html(page_html: str) -> str:
    """A temp file, not ``setHtml``: its data URL is capped at about 2 MB."""
    handle, path = tempfile.mkstemp(suffix=".html", prefix="contracting_print_")
    with os.fdopen(handle, "w", encoding="utf-8") as stream:
        stream.write(page_html)
    return path


def _remove_quietly(path: str | None) -> None:
    if path:
        try:
            os.remove(path)
        except OSError:
            pass


def _ask_pdf_path(parent: Any, document: PrintDocument) -> str:
    from PySide6.QtWidgets import QFileDialog

    path, _ = QFileDialog.getSaveFileName(parent, "حفظ PDF", f"{document.file_stem}.pdf", "PDF (*.pdf)")
    return path


def open_file(parent: Any, file_path: str) -> None:
    """A saved PDF / Excel opens straight away in its program (user, 2026-10-14)."""
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QDesktopServices
    from PySide6.QtWidgets import QMessageBox

    if not QDesktopServices.openUrl(QUrl.fromLocalFile(str(file_path))):
        QMessageBox.information(parent, "تم الحفظ", f"تم حفظ الملف، لكن تعذّر فتحه تلقائياً:\n{file_path}")


def _pdf_done(parent: Any, file_path: str, success: bool) -> None:
    from PySide6.QtWidgets import QMessageBox

    if success:
        open_file(parent, file_path)
    else:
        QMessageBox.warning(parent, "تعذّر التصدير", "حدث خطأ أثناء إنشاء ملف PDF.")


def _send_to_printer(parent: Any, view: Any, on_done: Any) -> bool:
    from PySide6.QtPrintSupport import QPrintDialog, QPrinter

    printer = QPrinter(QPrinter.HighResolution)
    printer.setPageLayout(page_layout())
    if QPrintDialog(printer, parent).exec() != QPrintDialog.Accepted:
        return False
    parent._contracting_printer = printer  # kept alive until Chromium finishes

    def _finished(success: bool) -> None:
        try:
            view.printFinished.disconnect(_finished)
        except (RuntimeError, TypeError):
            pass
        on_done(success)

    view.printFinished.connect(_finished)
    view.print(printer)
    return True


def _hidden_view(parent: Any, document: PrintDocument, on_ready: Any) -> None:
    """Load the document in a hidden view; call ``on_ready(view, cleanup)`` once laid out."""
    from PySide6.QtCore import QTimer, QUrl
    from PySide6.QtWidgets import QMessageBox
    from PySide6.QtWebEngineWidgets import QWebEngineView

    view = QWebEngineView(parent)
    view.setVisible(False)
    holder = parent.__dict__.setdefault("_contracting_print_views", [])
    holder.append(view)
    path = _write_temp_html(build_html(document))

    def cleanup() -> None:
        _remove_quietly(path)
        if view in holder:
            holder.remove(view)
        view.deleteLater()

    def _loaded(ok: bool) -> None:
        if not ok:
            QMessageBox.warning(parent, "الطباعة", "تعذّر تجهيز التقرير.")
            cleanup()
            return
        QTimer.singleShot(150, lambda: on_ready(view, cleanup))

    view.loadFinished.connect(_loaded)
    view.setUrl(QUrl.fromLocalFile(path))


def print_report(parent: Any, document: PrintDocument) -> None:
    """Straight to the printer (the print dialog, no preview)."""
    from PySide6.QtWidgets import QMessageBox

    def ready(view: Any, cleanup: Any) -> None:
        def done(success: bool) -> None:
            if not success:
                QMessageBox.warning(parent, "الطباعة", "تعذّرت طباعة التقرير.")
            cleanup()

        if not _send_to_printer(parent, view, done):
            cleanup()

    _hidden_view(parent, document, ready)


def export_pdf(parent: Any, document: PrintDocument) -> None:
    path = _ask_pdf_path(parent, document)
    if not path:
        return

    def ready(view: Any, cleanup: Any) -> None:
        page = view.page()

        def finished(file_path: str, success: bool) -> None:
            _pdf_done(parent, file_path, success)
            cleanup()

        page.pdfPrintingFinished.connect(finished)
        page.printToPdf(path, page_layout())

    _hidden_view(parent, document, ready)


def save_excel(parent: Any, document: PrintDocument) -> None:
    from PySide6.QtWidgets import QFileDialog, QMessageBox

    path, _ = QFileDialog.getSaveFileName(parent, "حفظ Excel", f"{document.file_stem}.xlsx", "Excel (*.xlsx)")
    if not path:
        return
    try:
        export_excel(document, path)
    except Exception as exc:  # noqa: BLE001 - shown to the user
        QMessageBox.warning(parent, "تعذّر التصدير", f"تعذّر حفظ ملف Excel:\n{exc}")
        return
    open_file(parent, path)


def preview_dialog(parent: Any, document: PrintDocument, can_print: bool, can_export: bool) -> Any:
    """«معاينة»: the pages as they will print, with طباعة / PDF / Excel / إغلاق."""
    from PySide6.QtCore import Qt, QUrl
    from PySide6.QtWidgets import QDialog, QHBoxLayout, QPushButton, QVBoxLayout
    from PySide6.QtWebEngineWidgets import QWebEngineView

    from app.ui.common.theme import GREEN, GREEN_DARK, _button_style

    dialog = QDialog(parent)
    dialog.setWindowTitle(f"معاينة — {document.title}")
    dialog.setLayoutDirection(Qt.RightToLeft)
    dialog.resize(1200, 860)
    root = QVBoxLayout(dialog)
    root.setContentsMargins(12, 12, 12, 12)
    root.setSpacing(10)
    bar = QHBoxLayout()
    bar.setSpacing(8)
    view = QWebEngineView(dialog)
    path = _write_temp_html(build_html(document))
    dialog.finished.connect(lambda *_: _remove_quietly(path))

    def _pdf() -> None:
        target = _ask_pdf_path(dialog, document)
        if not target:
            return
        page = view.page()

        def finished(file_path: str, success: bool) -> None:
            try:
                page.pdfPrintingFinished.disconnect(finished)
            except (RuntimeError, TypeError):
                pass
            _pdf_done(dialog, file_path, success)

        page.pdfPrintingFinished.connect(finished)
        page.printToPdf(target, page_layout())

    buttons = []
    for text, colors, action, allowed in (
        ("طباعة", (GREEN, GREEN_DARK), lambda: _send_to_printer(dialog, view, lambda _ok: None), can_print),
        ("تصدير PDF", ("#B91C1C", "#991B1B"), _pdf, can_export),
        ("تصدير Excel", ("#15803D", "#166534"), lambda: save_excel(dialog, document), can_export),
    ):
        button = QPushButton(text)
        button.setFixedHeight(36)
        button.setMinimumWidth(110)
        button.setStyleSheet(_button_style(*colors))
        button.setEnabled(allowed)
        button.clicked.connect(action)
        bar.addWidget(button)
        buttons.append(button)
    bar.addStretch(1)
    close = QPushButton("إغلاق")
    close.setFixedHeight(36)
    close.setMinimumWidth(90)
    close.setStyleSheet(_button_style("#374151", "#1F2937"))
    close.clicked.connect(dialog.reject)
    bar.addWidget(close)
    root.addLayout(bar)
    root.addWidget(view, 1)
    view.setUrl(QUrl.fromLocalFile(path))
    dialog.view, dialog.buttons = view, buttons
    return dialog


__all__ = [
    "DATE", "MONEY", "PCT", "TEXT", "SERIAL",
    "PrintColumn", "PrintDocument", "PrintSection",
    "build_html", "cell_text", "export_excel", "export_pdf", "label_span", "open_file", "page_layout",
    "preview_dialog", "print_report", "save_excel", "total_cells",
]
