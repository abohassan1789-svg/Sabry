"""طباعة تقرير كشف حساب الدفعة المقدمة — نموذج 3 «كشف لكل عقد بشريط السداد» (user, 2026-10-14).

Each contract prints as a block, as its window on the screen does: a band with the
contract and the share of its advance already used, the table of its extracts and
دفعات مقدمة, and beside it the contract's figures and the bar of what was paid back.

The user's rules:

* Each company prints on its own pages under نموذج 4's header (logo, name, address,
  السجل الضريبي from الشركات والمشاريع), repeated on every page with «صفحة س من ص».
* «تحديد أعمدة الجدول» decides what a block shows (the band's names and dates, the
  figures beside the table) and the columns of the summary table.
* After the blocks, a summary table of the company's contracts with its total; the
  last company's ends with «الإجمالي العام — كل الشركات».
* The bar is coloured as on the screen: greens for the extracts, blues for the دفعات
  مقدمة, grey for what is left, a red mark when a contract used more than its advance.
* Excel: one right-to-left sheet per company with the blocks one after the other and
  the summary table at the end, under the same header.

A block moves whole to the next page when it does not fit; only a block longer than a
page is split between its rows (its band then says «(تابع)»). The page, the preview,
the printer and the PDF are :mod:`app.reports.contracting_print`'s.
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.reports.contracting_print import (
    CELL_RULE,
    DATE,
    INK,
    MONEY,
    MUTED,
    FLOW_PAGINATE,
    PAGE_HEIGHT_PX,
    PCT,
    RED,
    SLATE,
    TOTAL_RULE,
    PrintColumn,
    _STYLE,
    _add_logo,
    _colgroup,
    _decimal,
    _esc,
    _excel_value,
    _filters_html,
    _header_html,
    _sheet_title,
    cell_text,
    company_lines,
    font_px,
    label_span,
    total_cells,
)

PAYMENT_LINE = "payment"
REST_COLOR = "#E5E7EB"
PAYMENT_BG = "#EFF6FF"
EXTRACT_COLOR, PAYMENT_COLOR = "#137A38", "#1D4ED8"  # the legend's two swatches

# The block's figures beside its table: (report key, caption); only the visible ones print.
FIGURES = (
    ("contract_value", "قيمة العقد"),
    ("advance_agreed", "الدفعة المقدمة"),
    ("advance_deducted", "المسدد من المستخلصات"),
    ("advance_paid_out", "المصروف من الدفعات"),
)
REMAINING = "remaining_advance"

# The block's table: the window's lines, without the names the band already carries.
LINE_COLUMNS = [
    PrintColumn("extract_date", "التاريخ", DATE, 78, nowrap=True),
    PrintColumn("extract_no", "رقم المستخلص / البيان"),
    PrintColumn("works_value", "إجمالي المستخلص", MONEY),
    PrintColumn("advance_payment_pct", "نسبة الدفعة المقدمة", PCT, 74),
    PrintColumn("advance_payment", "قيمة الدفعة المقدمة", MONEY),
]


@dataclass
class ContractBlock:
    """One contract: its report row, its lines (extracts and دفعات مقدمة, oldest first)
    and each line's colour on the bar."""

    contract: dict[str, Any]
    lines: list[dict[str, Any]]
    colors: list[str] = field(default_factory=list)


@dataclass
class AdvanceSection:
    """One company's pages."""

    company: dict[str, Any]
    blocks: list[ContractBlock]
    totals: dict[str, Any]

    @property
    def rows(self) -> list[dict[str, Any]]:
        return [block.contract for block in self.blocks]


@dataclass
class AdvanceDocument:
    title: str
    subtitle: str
    filters: list[tuple[str, str]]
    visible: tuple[str, ...]          # «تحديد أعمدة الجدول»: the report keys shown on the screen
    columns: list[PrintColumn]        # the summary table's columns
    sections: list[AdvanceSection]
    grand_totals: dict[str, Any] | None = None  # set when there is more than one company
    printed_by: str = ""
    printed_at: datetime.datetime = field(default_factory=datetime.datetime.now)
    count_noun: str = "عقود"

    @property
    def file_stem(self) -> str:
        return f"{self.title} - {self.printed_at:%Y-%m-%d}"

    def html(self) -> str:
        return build_html(self)

    def excel(self, path: str | Path) -> Path:
        return export_excel(self, path)


def _money(value: Any) -> str:
    return cell_text(PrintColumn("", "", MONEY), value)


def _pct(value: Any) -> str:
    return cell_text(PrintColumn("", "", PCT), value)


def used_pct(contract: dict[str, Any]) -> Decimal:
    """How much of the advance is used: deducted by the extracts and paid out."""
    agreed = _decimal(contract.get("advance_agreed"))
    used = _decimal(contract.get("advance_deducted")) + _decimal(contract.get("advance_paid_out"))
    return used / agreed * 100 if agreed > 0 else Decimal(0)


def bar_parts(block: ContractBlock) -> tuple[list[tuple[float, str]], float | None]:
    """The bar, right to left: (share of the bar, colour) per line and the grey rest, and
    where the advance ends (share from the right) when the lines took more than it."""
    agreed = _decimal(block.contract.get("advance_agreed"))
    amounts = [_decimal(line.get("advance_payment")) for line in block.lines]
    paid = sum(amounts, Decimal(0))
    whole = max(agreed, paid)
    if whole <= 0:
        return [], None
    parts = [(float(amount / whole), block.colors[index] if index < len(block.colors) else EXTRACT_COLOR)
             for index, amount in enumerate(amounts)]
    if agreed > paid:
        parts.append((float((agreed - paid) / whole), REST_COLOR))
    over = float(agreed / paid) if paid > agreed > 0 else None
    return parts, over


def band_parts(document: AdvanceDocument, contract: dict[str, Any]) -> list[str]:
    """The band after the contract: the names and date the screen shows."""
    parts = []
    for key in ("contractor_name", "project_name"):
        if key in document.visible and contract.get(key):
            parts.append(str(contract[key]))
    if "contract_date" in document.visible:
        parts.append(f"تاريخ العقد {cell_text(PrintColumn('', '', DATE), contract.get('contract_date'))}")
    return parts


def figure_cells(document: AdvanceDocument, contract: dict[str, Any]) -> list[tuple[str, Any]]:
    """(caption, amount) of the visible figures; the advance's caption carries its rate."""
    cells = []
    for key, caption in FIGURES:
        if key not in document.visible:
            continue
        if key == "advance_agreed" and "advance_payment_pct" in document.visible:
            caption = f"{caption} ({_pct(contract.get('advance_payment_pct'))})"
        cells.append((caption, contract.get(key)))
    return cells


def line_values(line: dict[str, Any]) -> dict[str, Any]:
    """A line's cells: a دفعة مقدمة has no extract total nor rate."""
    return {column.key: line.get(column.key) for column in LINE_COLUMNS}


def lines_caption(lines: list[dict[str, Any]]) -> str:
    extracts = sum(1 for line in lines if line.get("kind") != PAYMENT_LINE)
    payments = len(lines) - extracts
    parts = []
    if extracts:
        parts.append(f"{extracts} مستخلص" if extracts == 1 else f"{extracts} مستخلصات")
    if payments:
        parts.append(f"{payments} دفعة مقدمة" if payments == 1 else f"{payments} دفعات مقدمة")
    return "إجمالي المسدد — " + " + ".join(parts)


def lines_totals(lines: list[dict[str, Any]]) -> dict[str, Any]:
    """The total row; «إجمالي المستخلص» stays blank (``None``) when there is no extract."""
    extracts = [line for line in lines if line.get("kind") != PAYMENT_LINE]
    return {
        "works_value": sum((_decimal(line.get("works_value")) for line in extracts), Decimal(0)) if extracts else None,
        "advance_payment": sum((_decimal(line.get("advance_payment")) for line in lines), Decimal(0)),
    }


# -- HTML -------------------------------------------------------------------------------------

_EXTRA_STYLE = f"""
.legend {{ display:flex; gap:16px; align-items:center; font-size:10.5px; color:{MUTED}; }}
.legend i {{ display:inline-block; width:10px; height:10px; vertical-align:middle; margin-left:4px; }}
.blk {{ border:1.5px solid {SLATE}; border-radius:6px; overflow:hidden; display:flex; flex-direction:column; }}
.bh {{ display:flex; gap:16px; align-items:center; background:{SLATE}; color:#FFFFFF; padding:4px 12px; font-size:12px; }}
.bh b {{ font-size:13.5px; }}
.bh .used {{ margin-right:auto; white-space:nowrap; }}
.bb {{ display:grid; grid-template-columns:3fr 2fr; gap:12px; padding:7px 10px; align-items:start; }}
.bb th {{ background:#475569; border-color:#475569; padding:4px 3px; }}
.bb td {{ padding:4px 3px; }}
tr.pay td {{ background:{PAYMENT_BG}; color:#1E3A8A; }}
.side {{ display:flex; flex-direction:column; gap:6px; }}
.fg {{ display:grid; grid-template-columns:repeat(2, minmax(0, 1fr)); gap:4px 10px; font-size:10.5px; color:{MUTED}; }}
.fg b {{ display:block; font-size:12px; color:{INK}; direction:ltr; text-align:right; }}
.bar {{ position:relative; display:flex; height:12px; border:1px solid {SLATE}; border-radius:3px; overflow:hidden;
        background:{REST_COLOR}; }}
.bar span {{ display:block; height:100%; border-left:1px solid #FFFFFF; }}
.bar .over {{ position:absolute; top:-2px; bottom:-2px; width:3px; background:{RED}; }}
.rem {{ display:flex; justify-content:space-between; font-size:11.5px; }}
.rem b {{ direction:ltr; }}
.rem b.neg {{ color:{RED}; }}
.over-note {{ font-size:10.5px; color:{RED}; }}
.cap {{ font-size:12.5px; font-weight:700; display:flex; align-items:center; gap:10px; }}
.cap b {{ background:{SLATE}; color:#FFFFFF; padding:2px 14px; border-radius:4px; }}
"""

# Pages: the header, the filter strip and the legend (``.fixed``) repeat on every page. A
# contract's block moves whole to the next page when it does not fit; a block longer than a
# page, and the summary table (``.split``), are split between their rows, carrying their band
# and table head along. Then «س من ص» on every page.
_PAGINATE = FLOW_PAGINATE


def _legend_html() -> str:
    return ('<div class="legend fixed">'
            f'<span><i style="background:{EXTRACT_COLOR}"></i>مسدد من مستخلص</span>'
            f'<span><i style="background:{PAYMENT_COLOR}"></i>مصروف في دفعة مقدمة</span>'
            f'<span><i style="background:{REST_COLOR}; border:1px solid #94A3B8"></i>المتبقي</span>'
            f'<span><i style="background:{RED}; width:3px"></i>تجاوز الدفعة المقدمة</span></div>')


def _number_td(column: PrintColumn, value: Any) -> str:
    if value is None:
        return "<td></td>"
    text = cell_text(column, value)
    if column.kind not in (MONEY, PCT):
        return f'<td{" class=nw" if column.nowrap else ""}>{_esc(text)}</td>'
    negative = column.kind == MONEY and _decimal(value) < 0
    return f'<td class="num{" neg" if negative else ""}">{_esc(text)}</td>'


def _lines_table_html(block: ContractBlock) -> str:
    head = "".join(f"<th>{_esc(column.label)}</th>" for column in LINE_COLUMNS)
    body = []
    for line in block.lines:
        values = line_values(line)
        css = ' class="pay"' if line.get("kind") == PAYMENT_LINE else ""
        body.append(f"<tr{css}>" + "".join(_number_td(column, values[column.key]) for column in LINE_COLUMNS) + "</tr>")
    if block.lines:
        totals = lines_totals(block.lines)
        body.append(f'<tr class="total"><td colspan="2">{_esc(lines_caption(block.lines))}</td>'
                    f'<td class="num">{"" if totals["works_value"] is None else _money(totals["works_value"])}</td><td></td>'
                    f'<td class="num">{_money(totals["advance_payment"])}</td></tr>')
    else:
        body.append(f'<tr><td class="empty" colspan="{len(LINE_COLUMNS)}">'
                    "لا توجد مستخلصات ولا دفعات مقدمة على هذا العقد حتى الآن.</td></tr>")
    return (f'<table style="font-size:11px">{_colgroup(LINE_COLUMNS)}<thead><tr>{head}</tr></thead>'
            f"<tbody>{''.join(body)}</tbody></table>")


def _bar_html(block: ContractBlock) -> str:
    parts, over = bar_parts(block)
    spans = "".join(f'<span style="width:{share * 100:.3f}%; background:{color}"></span>' for share, color in parts)
    mark = f'<i class="over" style="right:calc({over * 100:.3f}% - 1px)"></i>' if over is not None else ""
    return f'<div class="bar">{spans}{mark}</div>'


def _side_html(document: AdvanceDocument, block: ContractBlock) -> str:
    contract = block.contract
    figures = "".join(f"<div>{_esc(caption)}<b>{_money(value)}</b></div>"
                      for caption, value in figure_cells(document, contract))
    parts = [f'<div class="fg">{figures}</div>' if figures else "", _bar_html(block)]
    remaining = _decimal(contract.get(REMAINING))
    if REMAINING in document.visible:
        css = ' class="neg"' if remaining < 0 else ""
        parts.append(f'<div class="rem"><span>المتبقي</span><b{css}>{_money(remaining)}</b></div>')
    if remaining < 0:
        parts.append(f'<div class="over-note">تجاوز الدفعة المقدمة بـ {_money(-remaining)}</div>')
    return '<div class="side">' + "".join(parts) + "</div>"


def _block_html(document: AdvanceDocument, block: ContractBlock) -> str:
    contract = block.contract
    band = "".join(f"<span>{_esc(part)}</span>" for part in band_parts(document, contract))
    return (f'<div class="blk"><div class="bh"><b class="cap-text">{_esc(contract.get("contract_label"))}</b>{band}'
            f'<span class="used">مستخدم {used_pct(contract):.0f}%</span></div>'
            f'<div class="bb">{_lines_table_html(block)}{_side_html(document, block)}</div></div>')


def _summary_html(document: AdvanceDocument, section: AdvanceSection, last: bool) -> str:
    columns = document.columns
    head = "".join(f"<th>{_esc(column.label)}</th>" for column in columns)
    body = []
    for serial, row in enumerate(section.rows, start=1):
        cells = []
        for column in columns:
            value = serial if column.key == "_serial" else row.get(column.key)
            cells.append(_number_td(column, value) if column.kind in (MONEY, PCT)
                         else f'<td{" class=nw" if column.nowrap else ""}>{_esc(cell_text(column, value))}</td>')
        body.append("<tr>" + "".join(cells) + "</tr>")
    span = label_span(columns)

    def total_row(label: str, totals: dict[str, Any], css: str) -> str:
        cells = "".join(f'<td class="num{" neg" if negative else ""}">{_esc(text)}</td>'
                        for text, negative in total_cells(columns, totals))
        return f'<tr class="{css}"><td colspan="{span}">{_esc(label)}</td>{cells}</tr>'

    if section.rows:
        body.append(total_row(f"الإجمالي — {len(section.rows)} {document.count_noun}", section.totals, "total"))
    else:
        body.append(f'<tr><td class="empty" colspan="{len(columns)}">لا توجد عقود</td></tr>')
    if last and document.grand_totals is not None:
        body.append(total_row("الإجمالي العام — كل الشركات", document.grand_totals, "grand"))
    return (f'<div class="split"><div class="cap"><b class="cap-text">ملخص العقود</b></div>'
            f'<table style="font-size:{font_px(columns)}px">{_colgroup(columns)}'
            f"<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table></div>")


def _section_html(document: AdvanceDocument, section: AdvanceSection, last: bool) -> str:
    parts = [_header_html(document, section, fixed=True), _filters_html(document, fixed=True), _legend_html()]
    parts += [_block_html(document, block) for block in section.blocks]
    parts.append(_summary_html(document, section, last))
    return '<div class="section">' + "".join(parts) + "</div>"


def build_html(document: AdvanceDocument) -> str:
    sections = "".join(_section_html(document, section, index == len(document.sections) - 1)
                       for index, section in enumerate(document.sections))
    script = _PAGINATE % {"limit": PAGE_HEIGHT_PX}
    return (
        '<!DOCTYPE html><html dir="rtl" lang="ar"><head><meta charset="utf-8">'
        f"<title>{_esc(document.title)}</title><style>{_STYLE}{_EXTRA_STYLE}</style></head>"
        f"<body>{sections}<script>{script}</script></body></html>"
    )


# -- Excel ------------------------------------------------------------------------------------


def export_excel(document: AdvanceDocument, path: str | Path) -> Path:
    """One right-to-left sheet per company: the header lines, each contract's block (band,
    figures, lines and their total) one after the other, then the summary table."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    thin, dark = Side(style="thin", color=CELL_RULE[1:]), Side(style="thin", color=TOTAL_RULE[1:])
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    total_grid = Border(left=dark, right=dark, top=dark, bottom=dark)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    head_fill, band_fill = PatternFill("solid", fgColor=SLATE[1:]), PatternFill("solid", fgColor=SLATE[1:])
    sub_head = PatternFill("solid", fgColor="475569")
    total_fill = PatternFill("solid", fgColor="E2E8F0")
    pay_fill = PatternFill("solid", fgColor=PAYMENT_BG[1:])
    white_bold = Font(bold=True, color="FFFFFF")
    money = PrintColumn("", "", MONEY)
    used: set[str] = set()
    columns = document.columns

    for index, section in enumerate(document.sections):
        company = section.company
        sheet = workbook.create_sheet(_sheet_title(company.get("company_name"), used))
        sheet.sheet_view.rightToLeft = True
        width = max(len(columns), len(LINE_COLUMNS), len(FIGURES) + 1)
        first_col = 1
        if company.get("logo"):
            first_col = 2
            _add_logo(sheet, company)
        lines = [(company.get("company_name") or "", Font(bold=True, size=16)),
                 *((line, Font(size=11, color=SLATE[1:])) for line in company_lines(company)),
                 (document.title, Font(bold=True, size=14)),
                 (document.subtitle, Font(size=11, color=SLATE[1:])),
                 ("   |   ".join(f"{caption}: {value}" for caption, value in document.filters), Font(size=11)),
                 (f"تاريخ الطباعة: {document.printed_at:%Y-%m-%d %H:%M}   |   المستخدم: {document.printed_by or '—'}"
                  f"   |   عدد العقود: {len(section.rows)}", Font(size=11, color=SLATE[1:]))]
        row_no = 1
        for text, font in lines:
            cell = sheet.cell(row=row_no, column=first_col, value=text)
            cell.font, cell.alignment = font, right
            if width > first_col:
                sheet.merge_cells(start_row=row_no, start_column=first_col, end_row=row_no, end_column=width)
            row_no += 1
        row_no += 1

        for block in section.blocks:
            contract = block.contract
            # the band
            text = "   ·   ".join([str(contract.get("contract_label") or "")] + band_parts(document, contract)
                                  + [f"مستخدم {used_pct(contract):.0f}%"])
            band = sheet.cell(row=row_no, column=1, value=text)
            for col in range(1, width + 1):
                sheet.cell(row=row_no, column=col).fill = band_fill
            band.font, band.alignment = white_bold, right
            sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=width)
            row_no += 1
            # the figures: captions over amounts
            figures = figure_cells(document, contract)
            if REMAINING in document.visible:
                figures.append(("المتبقي", contract.get(REMAINING)))
            if figures:
                for col, (caption, value) in enumerate(figures, start=1):
                    top, figure = sheet.cell(row=row_no, column=col, value=caption), sheet.cell(row=row_no + 1, column=col)
                    _excel_value(figure, money, value)
                    top.font = Font(size=10, color=MUTED[1:])
                    figure.font = Font(bold=True, color=RED[1:] if _decimal(value) < 0 else INK[1:])
                    for cell in (top, figure):
                        cell.alignment, cell.border = center, grid
                row_no += 2
            # the lines
            for col, column in enumerate(LINE_COLUMNS, start=1):
                cell = sheet.cell(row=row_no, column=col, value=column.label)
                cell.font, cell.fill, cell.alignment, cell.border = white_bold, sub_head, center, grid
            row_no += 1
            if block.lines:
                for line in block.lines:
                    values = line_values(line)
                    payment = line.get("kind") == PAYMENT_LINE
                    for col, column in enumerate(LINE_COLUMNS, start=1):
                        cell = sheet.cell(row=row_no, column=col)
                        if values[column.key] is not None:
                            _excel_value(cell, column, values[column.key])
                        cell.alignment, cell.border = center, grid
                        if payment:
                            cell.fill, cell.font = pay_fill, Font(color="1E3A8A")
                    row_no += 1
                totals = lines_totals(block.lines)
                sheet.cell(row=row_no, column=1, value=lines_caption(block.lines))
                for col, column in enumerate(LINE_COLUMNS, start=1):
                    cell = sheet.cell(row=row_no, column=col)
                    if totals.get(column.key) is not None:
                        _excel_value(cell, column, totals[column.key])
                    cell.fill, cell.font, cell.border, cell.alignment = total_fill, Font(bold=True), total_grid, center
                sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=2)
            else:
                sheet.cell(row=row_no, column=1, value="لا توجد مستخلصات ولا دفعات مقدمة على هذا العقد حتى الآن.")
                sheet.cell(row=row_no, column=1).alignment = center
                sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=len(LINE_COLUMNS))
            row_no += 2

        # the summary table
        sheet.cell(row=row_no, column=1, value="ملخص العقود").font = Font(bold=True, size=12)
        row_no += 1
        for col, column in enumerate(columns, start=1):
            cell = sheet.cell(row=row_no, column=col, value=column.label)
            cell.font, cell.fill, cell.alignment, cell.border = white_bold, head_fill, center, grid
        sheet.row_dimensions[row_no].height = 32
        row_no += 1
        for serial, row in enumerate(section.rows, start=1):
            for col, column in enumerate(columns, start=1):
                value = serial if column.key == "_serial" else row.get(column.key)
                cell = sheet.cell(row=row_no, column=col)
                _excel_value(cell, column, value)
                cell.alignment, cell.border = center, grid
                if column.kind == MONEY and _decimal(value) < 0:
                    cell.font = Font(color=RED[1:])
            row_no += 1
        span = label_span(columns)

        def total_row(row_no: int, label: str, totals: dict[str, Any], fill: Any, font: Any) -> int:
            sheet.cell(row=row_no, column=1, value=label)
            for col, column in enumerate(columns, start=1):
                cell = sheet.cell(row=row_no, column=col)
                if col > span and column.kind == MONEY and column.key in totals:
                    _excel_value(cell, column, totals[column.key])
                cell.fill, cell.font, cell.border, cell.alignment = fill, font, total_grid, center
            if span > 1:
                sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=span)
            return row_no + 1

        if section.rows:
            row_no = total_row(row_no, f"الإجمالي — {len(section.rows)} {document.count_noun}", section.totals,
                               total_fill, Font(bold=True))
        if index == len(document.sections) - 1 and document.grand_totals is not None:
            total_row(row_no, "الإجمالي العام — كل الشركات", document.grand_totals, head_fill, white_bold)

        for col in range(1, width + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


__all__ = [
    "FIGURES", "LINE_COLUMNS", "PAYMENT_LINE", "REST_COLOR",
    "AdvanceDocument", "AdvanceSection", "ContractBlock",
    "band_parts", "bar_parts", "build_html", "export_excel", "figure_cells", "line_values", "lines_caption",
    "lines_totals", "used_pct",
]
